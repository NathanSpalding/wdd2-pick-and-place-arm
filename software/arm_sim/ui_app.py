"""
WDD2 pick-and-place arm simulator: user interface
=================================================

Run:   python ui_app.py
Then open http://localhost:8080 in your browser.
Stop:  press Ctrl+C in the terminal.

Folder layout (inside arm_sim):
  ui_app.py
  WDD2_URDF_Export_description/
      urdf/WDD2_arm.urdf
      meshes/*.stl

Tabs
  Control   - status, type in X, Y, Z (mm) and the gripper angle, grab/release
  Sequence  - build (presets or custom points), run over several trials,
              save and load step-by-step tasks
  Joints    - joint angles, commanded vs measured angles (AS5600), speed limits
  Currents  - model servo currents, and model vs measured gripper current
              (the INA219 measures the gripper servo only)
  Report    - cycle report after each trial, and the one-file data export
  Hardware  - serial connection to the ESP, and the message format
  Setup     - display options (incl. joint frames), floor check and the
              workspace map

All positions shown in the interface are in millimetres, angles in degrees.
Joint angles: 0 = arm stretched out horizontally; positive shoulder, elbow and
wrist angles lift the arm.

The EMERGENCY STOP button at the top is always visible. The physical E-stop on
the arm also triggers it (the ESP sends "E,1").

Until the ESP sends real readings, the "measured" angles and gripper current
are fictitious (made from the simulated values with a lag, an offset and noise).

Everything you might want to change is in SETTINGS below.
"""

import csv
import json
import threading
import time
from collections import deque
from datetime import datetime
from pathlib import Path

import numpy as np
import pinocchio as pin
import pink
import plotly.graph_objects as go
import viser
from pink import solve_ik
from pink.barriers import PositionBarrier
from pink.tasks import FrameTask
import yourdfpy
from viser.extras import ViserUrdf

try:
    import serial
    import serial.tools.list_ports
    HAVE_SERIAL = True
except ImportError:
    HAVE_SERIAL = False

# =============================================================================
# SETTINGS  (change these when you switch to your own arm)
# =============================================================================
HERE = Path(__file__).resolve().parent
URDF_PATH = HERE / "WDD2_URDF_Export_description" / "urdf" / "WDD2_arm.urdf"
TCP_FRAME = "tcp"                # gripping point, 133 mm beyond the wrist joint

# Arm joints, in the order of the URDF and of the J / A / I messages
JOINT_NAMES = ["Base", "Shoulder", "Elbow", "Wrist"]
NUM_JOINTS = len(JOINT_NAMES)
SHOULDER_JOINT = "shoulder"      # URDF joint name, used for the reach check

# Servo data per joint (base DS3240-270, shoulder DS5180, elbow DS3240,
# wrist DS3225). Torque constant = stall torque / stall current.
STALL_TORQUE = np.array([4.36, 9.61, 4.36, 2.45])  # Nm
STALL_CURRENT = np.array([3.9, 6.2, 3.9, 2.3])     # A (= current limit per servo)
SERVO_VOLTAGE = np.array([6.7, 8.0, 6.7, 6.5])     # V, rail of each servo
KT = STALL_TORQUE / STALL_CURRENT                  # Nm per A
CURRENT_LIMIT = STALL_CURRENT
IDLE_CURRENT = 0.15              # A per servo when not loaded

# Gripper servo (MG996R, 6.5 V rail). It is the only servo whose current is
# measured on the arm (INA219, sent by the ESP as I,<amps>).
GRIPPER_NAME = "Gripper"
GRIP_IDLE_CURRENT = 0.009        # A, jaws open (measured, bring-up 9 Oct 2026)
GRIP_HOLD_CURRENT = 0.11         # A, holding the 0.5 kg cube at 950 us (measured, bring-up 9 Oct 2026)
GRIP_STALL_CURRENT = 2.5         # A
GRIP_VOLTAGE = 6.5               # V

# Current channels: the four joints (model only) plus the gripper
CUR_NAMES = JOINT_NAMES + [GRIPPER_NAME]
CUR_LIMIT = np.append(CURRENT_LIMIT, GRIP_STALL_CURRENT)
CUR_VOLTAGE = np.append(SERVO_VOLTAGE, GRIP_VOLTAGE)

# Floor protection: the tip and the wrist are never allowed below these heights
TIP_MIN_Z_MM = 5.0               # gripper tip (TCP)
WRIST_MIN_Z_MM = 20.0            # wrist joint (the wrist servo body sits around it)
WRIST_FRAME = "gripper_link"     # frame at the wrist joint
# Moves follow an arc around the base axis (radius, angle and height change
# together) instead of a straight line through the middle. The IK chases a point
# this far ahead of the tip along that arc.
LOOKAHEAD_MM = 40.0

MAX_JOINT_SPEED = 1.5            # rad/s, starting value (editable in Joints tab)
MAX_JOINT_ACCEL = 8.0            # rad/s^2, starting value (editable in Joints tab)
IN_POSITION_MM = 5.0             # "arrived" when the tip is this close (mm)

# Gripper angle (tilt) relative to horizontal: -90 = pointing straight down,
# 0 = level. With 4 joints the arm can set the tip position plus this one angle.
# Auto mode: the gripper points down where it can (about 280-360 mm from the base
# at cube height) and tilts toward level further out (level reaches ~500 mm).
GRIPPER_ANGLE_DEG = -90.0        # used in Fixed mode
AUTO_ANGLES_DEG = list(range(-90, 1, 10))   # tried in order: steepest first
AUTO, FIXED = "Auto (by reach)", "Fixed"

# Starting pose in degrees (base, shoulder, elbow, wrist): gripper pointing down
HOME_DEG = [0.0, 72.0, -82.0, -80.0]
BASE_AIM_AT_TIP = True        # real arm: tool point in line with the base axis (see j_q below)

CUBE_SIZE_MM = 40.0              # edge length
CUBE_MASS = 0.5                  # kg
CUBE_HOME_MM = (300.0, 0.0)      # where the cube starts (x, y)

# Named positions of the gripper tip (x, y, z) in millimetres. "Home" is added
# automatically from the starting pose.
# Note: with the gripper pointing down, the wrist's -85 deg limit keeps the tip
# below about shoulder height + 40 mm, so the "above" points are kept low.
PRESETS_MM = {
    "Above pick": (290, 0, 70),
    "Pick": (290, 0, CUBE_SIZE_MM / 2),
    "Above place": (205, 205, 70),
    "Place": (205, 205, CUBE_SIZE_MM / 2),
}

BAUD_RATES = ["115200", "230400", "921600", "57600", "9600"]
DEMO_PORT = "Simulated ESP (demo)"
MEASURED_TIMEOUT = 0.5           # s: ESP readings older than this are ignored

WORKSPACE_STEP_MM = 20.0         # grid spacing of the workspace map
WORKSPACE_RADIUS_MM = 560.0      # how far out the map scans

# Fictitious "measured" joint angles (used until the ESP sends A messages)
FAKE_ANGLE_LAG = 0.12            # s, servo response lag
FAKE_ANGLE_NOISE_DEG = 0.3       # encoder noise
FAKE_SAG_DEG_PER_NM = 0.8        # extra droop under gravity load

DT = 0.02                        # control loop period (50 Hz)
LOG_EVERY = 1                    # log a CSV row every loop (50 Hz; was 5 = 10 Hz)
BRAND = (230, 120, 40)           # accent colour for the interface
JOINT_COLOURS = ["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728", "#9467bd"]  # + gripper

LOG_DIR = HERE / "logs"

# ---- Internal units are metres and radians. Do not edit below. ----
MM = 1000.0
CUBE_SIZE = CUBE_SIZE_MM / MM
CUBE_HOME = (CUBE_HOME_MM[0] / MM, CUBE_HOME_MM[1] / MM)
PRESETS = {k: tuple(v / MM for v in p) for k, p in PRESETS_MM.items()}
WORKSPACE_STEP = WORKSPACE_STEP_MM / MM
WORKSPACE_RADIUS = WORKSPACE_RADIUS_MM / MM
CUSTOM = "Custom"
SEQ_DIR = HERE / "saved_sequences"

if not URDF_PATH.exists():
    raise SystemExit(
        f"\nCould not find the arm model:\n  {URDF_PATH}\n"
        "Copy the whole WDD2_URDF_Export_description folder (with urdf/WDD2_arm.urdf "
        "and meshes/) into the same folder as ui_app.py.\n")

# =============================================================================
# ROBOT MATHS (Pinocchio + Pink)
# =============================================================================
model = pin.buildModelFromUrdf(str(URDF_PATH))
TCP_ID = model.getFrameId(TCP_FRAME)
if TCP_ID >= model.nframes:
    raise SystemExit(f"The URDF has no '{TCP_FRAME}' link (the gripper tip).")

data = model.createData()
data_dyn = model.createData()

q0 = np.radians(HOME_DEG)
configuration = pink.Configuration(model, data, q0)

# With 4 joints the arm controls the tip position plus ONE angle: the gripper
# tilt (rotation about the local X axis, the axis of shoulder, elbow and wrist).
# So the orientation cost is only on that axis.
task = FrameTask(TCP_FRAME, position_cost=1.0, orientation_cost=[0.03, 0.0, 0.0])

LOWER = model.lowerPositionLimit[:NUM_JOINTS]
UPPER = model.upperPositionLimit[:NUM_JOINTS]
VEL_LIMIT = model.velocityLimit[:NUM_JOINTS]

# Floor barriers: Pink keeps these frames above the given heights at all times
TIP_MIN_Z = TIP_MIN_Z_MM / 1000.0
WRIST_MIN_Z = WRIST_MIN_Z_MM / 1000.0
WRIST_ID = model.getFrameId(WRIST_FRAME)
floor_barriers = [
    PositionBarrier(TCP_FRAME, indices=[2], p_min=np.array([TIP_MIN_Z]),
                    gain=np.array([10.0]), safe_displacement_gain=0.0),
    PositionBarrier(WRIST_FRAME, indices=[2], p_min=np.array([WRIST_MIN_Z]),
                    gain=np.array([10.0]), safe_displacement_gain=0.0),
]


def above_floor(d):
    """True if tip and wrist are above their minimum heights (d: updated data)."""
    return (d.oMf[TCP_ID].translation[2] >= TIP_MIN_Z - 1e-3
            and d.oMf[WRIST_ID].translation[2] >= WRIST_MIN_Z - 1e-3)


def to_cyl(p):
    return np.array([np.hypot(p[0], p[1]), np.arctan2(p[1], p[0]), p[2]])


def from_cyl(c):
    return np.array([c[0] * np.cos(c[1]), c[0] * np.sin(c[1]), c[2]])


def arc_point(a, b, s):
    """Point a fraction s along the arc from a to b around the base axis."""
    ca, cb = to_cyl(a), to_cyl(b)
    dth = (cb[1] - ca[1] + np.pi) % (2 * np.pi) - np.pi     # shortest way round
    return from_cyl([ca[0] + s * (cb[0] - ca[0]), ca[1] + s * dth,
                     ca[2] + s * (cb[2] - ca[2])])


def arc_length(a, b):
    ca, cb = to_cyl(a), to_cyl(b)
    dth = (cb[1] - ca[1] + np.pi) % (2 * np.pi) - np.pi
    return float(np.sqrt((cb[0] - ca[0]) ** 2 + (0.5 * (ca[0] + cb[0]) * dth) ** 2
                         + (cb[2] - ca[2]) ** 2))


def arc_carrot(tip, goal):
    """The point the IK aims for: LOOKAHEAD_MM ahead of the tip along the arc."""
    L = arc_length(tip, goal)
    if L <= LOOKAHEAD_MM / 1000.0:
        return np.asarray(goal, dtype=float)
    return arc_point(tip, goal, LOOKAHEAD_MM / 1000.0 / L)


def gripper_angle(q):
    """Gripper tilt relative to horizontal (rad): shoulder + elbow + wrist,
    because their axes are parallel."""
    return float(q[1] + q[2] + q[3])


def tilt_target(q, p, angle):
    """Target pose for Pink: position p, current heading, gripper tilt = angle."""
    R = configuration.get_transform_frame_to_world(TCP_FRAME).rotation
    R_goal = R @ pin.utils.rotate("x", angle - gripper_angle(q))
    return pin.SE3(R_goal, np.asarray(p, dtype=float))


def tip_position():
    return configuration.get_transform_frame_to_world(TCP_FRAME).translation.copy()


STEP_DT = [DT]   # real time since the last loop (s), so motion keeps its speed if a loop runs late


def advance(qd):
    """Move the simulated arm by qd for one time step, never past a joint limit."""
    q = pin.integrate(model, configuration.q, qd * STEP_DT[0])
    configuration.update(np.clip(q, model.lowerPositionLimit, model.upperPositionLimit))


# Shoulder position and reach, worked out from the model
pin.forwardKinematics(model, data_dyn, np.zeros(model.nq))
SHOULDER = data_dyn.oMi[model.getJointId(SHOULDER_JOINT)].translation.copy()
pin.framesForwardKinematics(model, data_dyn, np.zeros(model.nq))
MAX_REACH = float(np.linalg.norm(data_dyn.oMf[TCP_ID].translation - SHOULDER))

# The shoulder turns with the base: it sits SHOULDER_R from the base axis.
SHOULDER_R = float(np.hypot(SHOULDER[0], SHOULDER[1]))


def beyond_reach(p, margin=0.0):
    """Quick test: True if p is further than the arm can stretch, for any base
    angle (distance to the circle the shoulder moves on)."""
    horiz = max(0.0, np.hypot(p[0], p[1]) - SHOULDER_R)
    return np.hypot(horiz, p[2] - SHOULDER[2]) > MAX_REACH + margin


HOME = tip_position()
PRESETS = {"Home": tuple(HOME), **PRESETS}

# =============================================================================
# VIEWER AND 3D SCENE
# =============================================================================
server = viser.ViserServer()
server.scene.set_up_direction("+z")

urdf = yourdfpy.URDF.load(str(URDF_PATH))
viser_urdf = ViserUrdf(server, urdf_or_path=urdf)
viewer_joint_names = list(viser_urdf.get_actuated_joint_limits().keys())

# Floor check: the lowest point of the model with all joints at 0 should be z = 0
try:
    urdf.update_cfg(np.zeros(len(viewer_joint_names)))
    FLOOR_GAP_MM = float(urdf.scene.bounds[0][2]) * MM
except Exception:
    FLOOR_GAP_MM = 0.0
if abs(FLOOR_GAP_MM) > 2.0:
    FLOOR_MSG = (f"⚠️ The base bottom is at Z = {FLOOR_GAP_MM:.0f} mm, not on the floor. "
                 f"In WDD2_arm.urdf set the last xyz value of `world_to_base` to "
                 f"**{-FLOOR_GAP_MM / MM:.4f}**, then restart.")
    print("\n" + FLOOR_MSG.replace("**", "").replace("`", "") + "\n")
else:
    FLOOR_MSG = "✅ The base sits on the floor (Z = 0)."


def show_robot(q):
    cfg = [q[model.joints[model.getJointId(n)].idx_q] for n in viewer_joint_names]
    viser_urdf.update_cfg(np.array(cfg))


ground = server.scene.add_grid(
    "/ground", width=1.6, height=1.6, plane="xy",
    cell_size=0.05, section_size=0.25, plane_opacity=1.0,
    position=(0.0, 0.0, -0.001),
)
server.scene.add_frame("/origin", axes_length=0.08, axes_radius=0.003)

# Joint frames (Setup tab > Show joint frames): one per URDF link frame, which
# sits at the joint that drives it, plus the gripping point (TCP).
FRAME_LINKS = [("turntable_link", "Base joint"), ("upper_arm_link", "Shoulder joint"),
               ("forearm_link", "Elbow joint"), ("gripper_link", "Wrist joint"),
               (TCP_FRAME, "TCP")]
frame_nodes = []
for _link, _text in FRAME_LINKS:
    _f = server.scene.add_frame(f"/frames/{_link}", axes_length=0.06,
                                axes_radius=0.0025, visible=False)
    server.scene.add_label(f"/frames/{_link}/label", text=_text,
                           position=(0.0, 0.0, 0.0))
    frame_nodes.append((_link, _f))


def update_frames():
    """Move the joint frames to the current pose (only while they are shown)."""
    for link, node in frame_nodes:
        T = configuration.get_transform_frame_to_world(link)
        x, y, z, w = pin.Quaternion(T.rotation).coeffs()
        node.wxyz = (float(w), float(x), float(y), float(z))
        node.position = tuple(float(v) for v in T.translation)

cube = server.scene.add_box(
    "/cube", color=(230, 150, 40),
    dimensions=(CUBE_SIZE, CUBE_SIZE, CUBE_SIZE),
    position=(CUBE_HOME[0], CUBE_HOME[1], CUBE_SIZE / 2),
)
cube_label = server.scene.add_label(
    "/cube/label", f"{CUBE_MASS * 1000:.0f} g", position=(0.0, 0.0, CUBE_SIZE),
    anchor="bottom-center",
)

target = server.scene.add_transform_controls(
    "/target", scale=0.1, position=tuple(HOME), disable_rotations=True,
)
target_ball = server.scene.add_icosphere(
    "/target_ball", radius=0.012, color=(60, 200, 90), opacity=0.7,
    position=tuple(HOME),
)
trail_points = deque(maxlen=150)
trail = None
workspace_cloud = None


def apply_scene_theme(dark):
    if dark:
        ground.plane_color = (45, 48, 55)
        ground.cell_color = (70, 74, 82)
        ground.section_color = (110, 115, 125)
    else:
        ground.plane_color = (228, 231, 235)
        ground.cell_color = (200, 204, 210)
        ground.section_color = (150, 155, 165)


apply_scene_theme(False)


def reset_camera(client):
    client.camera.position = (0.75, -0.65, 0.5)
    client.camera.look_at = (0.2, 0.0, 0.1)


@server.on_client_connect
def _(client):
    reset_camera(client)


def notify(title, body, color=None):
    for client in server.get_clients().values():
        client.add_notification(title, body, auto_close_seconds=4.0, color=color)


# =============================================================================
# SHARED STATE
# =============================================================================
state = {
    "mode": "Manual",          # Manual / Sequence
    "estop": False,
    "holding": False,
    "reach": "ok",             # ok / moving / out
    "stall_time": 0.0,
    "last_target": np.array(HOME),
    "warned_target": None,
    "prev_qd": np.zeros(model.nv),
    "stopping": False,
}
seq = {"steps": [], "index": 0, "started": False, "t0": 0.0,
       "arrived_at": None, "trial": 1, "trials_total": 1, "first_report": 0}
seq_nodes = []                 # 3D markers (ball, label) for the move steps
trials_summary = ""
hw = {"mode": "offline",       # offline / demo / serial
      "ser": None, "port": "", "measured": None, "measured_time": 0.0,
      "angles": None, "angles_time": 0.0,     # AS5600 readings (rad)
      "estop_pressed": False,                  # physical E-stop (E,1 / E,0)
      "estop_event": False,                    # set by the reader thread
      "events": [],                            # K,... / F,... lines for the main loop
      "info": "",                              # last '#' line from the ESP (status, J stats)
      "j_q": None, "j_qd": None, "j_time": 0.0,  # latest target (deg) and speed (deg/s) for the J sender
      "sent": 0, "received": 0, "last_rx": "", "error": ""}
hw["wp_id"] = 0          # id of the last waypoint (W,...,id) sent to the ESP
hw["wp_done"] = None     # id from the last Y,id reply (move finished on the ESP)
cycle = {"active": False, "kind": "", "t0": 0.0, "samples": []}
reports = []

history = deque(maxlen=300)          # (t, model currents, measured gripper) for the plot
angle_history = deque(maxlen=300)    # (t, commanded, measured) joint angles
angle_window = deque(maxlen=250)     # last 5 s of (commanded, measured)
compare_window = deque(maxlen=250)   # last 5 s of gripper (model, measured)
log_rows = []                        # for CSV export
log_start = time.time()
preset_nodes = {}

# Fictitious "measured" gripper current: gain error, offset, lag and noise.
_rng = np.random.default_rng(7)
FAKE_GAIN = 1.0 + _rng.uniform(-0.12, 0.15)
FAKE_OFFSET = _rng.uniform(0.03, 0.12)
FAKE_LAG = 0.08                      # s
fake_state = 0.0
FAKE_ANGLE_OFFSET = np.radians(_rng.uniform(-1.5, 1.5, NUM_JOINTS))  # calibration error
fake_angles = q0[:NUM_JOINTS].copy()

# =============================================================================
# GUI: always-visible header
# =============================================================================
server.gui.configure_theme(
    brand_color=BRAND, control_width="large", show_share_button=False,
)
server.gui.add_markdown("### WDD2 pick-and-place arm")
estop_btn = server.gui.add_button("EMERGENCY STOP", color="red",
                                  icon=viser.Icon.ALERT_OCTAGON)
estop_banner = server.gui.add_markdown("", visible=False)
estop_reset = server.gui.add_button("Reset emergency stop", color="yellow",
                                    icon=viser.Icon.RESTORE, visible=False)
mode_md = server.gui.add_markdown("")

tabs = server.gui.add_tab_group()

# ---------------- Control tab ----------------
with tabs.add_tab("Control", icon=viser.Icon.TARGET):
    with server.gui.add_folder("Status"):
        reach_md = server.gui.add_markdown("")
        tip_txt = server.gui.add_text("Tip position (mm)", initial_value="",
                                      disabled=True)
        grip_txt = server.gui.add_text("Gripper angle", initial_value="",
                                       disabled=True)

    with server.gui.add_folder("Move gripper tip"):
        x_in = server.gui.add_number("X (mm)", initial_value=round(HOME[0] * MM), step=1)
        y_in = server.gui.add_number("Y (mm)", initial_value=round(HOME[1] * MM), step=1)
        z_in = server.gui.add_number("Z (mm)", initial_value=round(HOME[2] * MM),
                                     step=1, min=0)
        grip_mode = server.gui.add_dropdown(
            "Gripper angle", [AUTO, FIXED], initial_value=AUTO,
            hint="Auto: point straight down where possible, tilt toward level for "
                 "points further out. Fixed: always use the angle below. "
                 "Used for every move, also in sequences.")
        grip_angle = server.gui.add_slider(
            "Fixed angle (°)", min=-90, max=30, step=5,
            initial_value=GRIPPER_ANGLE_DEG, visible=False,
            hint="-90 = pointing straight down, 0 = level")
        go_btn = server.gui.add_button("Go to position", icon=viser.Icon.CURRENT_LOCATION)
        speed = server.gui.add_slider(
            "Speed (%)", min=10, max=100, step=5, initial_value=60,
            hint="Percentage of the max joint speed set in the Joints tab")

    with server.gui.add_folder(f"Cube ({CUBE_MASS * 1000:.0f} g)"):
        cube_md = server.gui.add_markdown("On the ground")
        grab_btn = server.gui.add_button("Grab cube", color="green",
                                         icon=viser.Icon.HAND_GRAB)
        release_btn = server.gui.add_button("Release cube", icon=viser.Icon.ARROW_DOWN)
        cube_x = server.gui.add_number("Cube X (mm)", initial_value=CUBE_HOME_MM[0], step=1)
        cube_y = server.gui.add_number("Cube Y (mm)", initial_value=CUBE_HOME_MM[1], step=1)
        place_btn = server.gui.add_button("Put cube here", icon=viser.Icon.CUBE)

# ---------------- Sequence tab ----------------
STEP_TYPES = ["Move to position", "Grab cube", "Release cube", "Wait 1 s"]
with tabs.add_tab("Sequence", icon=viser.Icon.LIST):
    with server.gui.add_folder("Add step"):
        step_type = server.gui.add_dropdown("Step", STEP_TYPES)
        step_preset = server.gui.add_dropdown("Position", list(PRESETS.keys()) + [CUSTOM])
        cx_in = server.gui.add_number("Custom X (mm)", initial_value=round(HOME[0] * MM),
                                      step=1, visible=False)
        cy_in = server.gui.add_number("Custom Y (mm)", initial_value=round(HOME[1] * MM),
                                      step=1, visible=False)
        cz_in = server.gui.add_number("Custom Z (mm)", initial_value=round(HOME[2] * MM),
                                      step=1, visible=False)
        copy_btn = server.gui.add_button("Copy from 3D handle", icon=viser.Icon.TARGET,
                                         visible=False)
        check_md = server.gui.add_markdown("", visible=False)
        step_add = server.gui.add_button("Add step", icon=viser.Icon.PLUS)
        step_undo = server.gui.add_button("Undo last step", icon=viser.Icon.ARROW_BACK_UP)
        step_clear = server.gui.add_button("Clear all", icon=viser.Icon.TRASH)
        step_demo = server.gui.add_button("Load pick-and-place demo",
                                          icon=viser.Icon.PLAYLIST_ADD)
        with server.gui.add_folder("Save custom point as preset", expand_by_default=False):
            preset_name = server.gui.add_text("Preset name", initial_value="My point")
            preset_save = server.gui.add_button("Save as preset",
                                                icon=viser.Icon.DEVICE_FLOPPY)

    with server.gui.add_folder("Run"):
        trials_in = server.gui.add_number("Number of trials", initial_value=1,
                                          min=1, max=100, step=1)
        reset_cube_cb = server.gui.add_checkbox(
            "Reset cube before each trial", initial_value=True,
            hint="Put the cube back at Cube X, Y (Control tab) before every trial")
        seq_run = server.gui.add_button("Run sequence", color="green",
                                        icon=viser.Icon.PLAYER_PLAY)
        seq_stop = server.gui.add_button(
            "Stop", icon=viser.Icon.PLAYER_STOP,
            hint="Stops any motion (sequence or manual) and holds position")
        seq_progress = server.gui.add_progress_bar(0.0, color=BRAND)
        trial_md = server.gui.add_markdown("")
        seq_md = server.gui.add_markdown("*No steps yet*")

    with server.gui.add_folder("Saved sequences"):
        server.gui.add_markdown("Saves the steps and your presets together.")
        seq_name = server.gui.add_text("Sequence name", initial_value="pick_and_place")
        seq_save_btn = server.gui.add_button("Save sequence", color="green",
                                             icon=viser.Icon.DEVICE_FLOPPY)
        seq_saved_dd = server.gui.add_dropdown("Saved sequences", ["(none)"])
        seq_refresh_btn = server.gui.add_button("Refresh list", icon=viser.Icon.REFRESH)
        seq_load_btn = server.gui.add_button("Load sequence", icon=viser.Icon.FOLDER_OPEN)
        seq_delete_btn = server.gui.add_button("Delete saved sequence",
                                               icon=viser.Icon.TRASH)

# ---------------- Joints tab ----------------
with tabs.add_tab("Joints", icon=viser.Icon.ROBOT):
    joints_md = server.gui.add_markdown("")
    with server.gui.add_folder("Commanded vs measured angle"):
        angle_source_md = server.gui.add_markdown("")
        angle_mode = server.gui.add_dropdown(
            "Show", ["Both", "Error (meas - cmd)"])
        angle_joint = server.gui.add_dropdown("Joint ", ["All"] + JOINT_NAMES)
        angle_plot = server.gui.add_plotly(figure=go.Figure(), aspect=1.3)
        server.gui.add_markdown("*Solid = commanded, dotted = measured (AS5600).*")
        angles_md = server.gui.add_markdown("")
    with server.gui.add_folder("Speed and acceleration limits"):
        max_speed = server.gui.add_number("Max joint speed (rad/s)",
                                          initial_value=MAX_JOINT_SPEED, step=0.1, min=0.1)
        max_accel = server.gui.add_number("Max joint accel (rad/s²)",
                                          initial_value=MAX_JOINT_ACCEL, step=0.5, min=0.5)
        limits_md = server.gui.add_markdown("")

# ---------------- Currents tab ----------------
SHOW_MODEL = "Model currents (all servos)"
SHOW_GRIP = "Gripper: model vs measured"
SHOW_GRIP_ERR = "Gripper error (meas - model)"
with tabs.add_tab("Currents", icon=viser.Icon.CHART_LINE):
    source_md = server.gui.add_markdown("")
    plot_mode = server.gui.add_dropdown(
        "Show", [SHOW_MODEL, SHOW_GRIP, SHOW_GRIP_ERR],
        hint="Joint currents come from the arm model only; the gripper current "
             "is also measured (INA219)")
    plot_joint = server.gui.add_dropdown(
        "Servo", ["All"] + CUR_NAMES, hint="Used in the model view")
    plot = server.gui.add_plotly(figure=go.Figure(), aspect=1.3)
    server.gui.add_markdown("*Solid = model, dotted = measured (gripper only).*")
    currents_md = server.gui.add_markdown("Waiting for data...")
    with server.gui.add_folder("Data log"):
        recording = server.gui.add_checkbox("Recording", initial_value=True)
        log_txt = server.gui.add_text("Samples", initial_value="0", disabled=True)
        server.gui.add_markdown("*Export is in the Report tab.*")
        clear_log_btn = server.gui.add_button("Clear log", icon=viser.Icon.TRASH)

# ---------------- Report tab ----------------
with tabs.add_tab("Report", icon=viser.Icon.REPORT_ANALYTICS):
    server.gui.add_markdown("A report is made after each sequence trial "
                            "(or when a run is stopped).")
    report_md = server.gui.add_markdown("*No cycles run yet*")
    with server.gui.add_folder("Export"):
        server.gui.add_markdown(
            "One Excel file with three sheets: **Summary**, **Cycle reports** "
            "and **Data log** (every sample from the Currents tab).")
        export_all_btn = server.gui.add_button("Export all data", color="green",
                                               icon=viser.Icon.FILE_SPREADSHEET)
    with server.gui.add_folder("All cycles this session"):
        report_hist_md = server.gui.add_markdown("*None yet*")
        report_clear = server.gui.add_button("Clear reports", icon=viser.Icon.TRASH)

# ---------------- Hardware tab ----------------
with tabs.add_tab("Hardware", icon=viser.Icon.CPU):
    hw_status_md = server.gui.add_markdown("")
    with server.gui.add_folder("Connection"):
        port_dd = server.gui.add_dropdown("Port", [DEMO_PORT])
        baud_dd = server.gui.add_dropdown("Baud rate", BAUD_RATES)
        port_refresh = server.gui.add_button("Refresh ports", icon=viser.Icon.REFRESH)
        connect_btn = server.gui.add_button("Connect", color="green",
                                            icon=viser.Icon.PLUG_CONNECTED)
        disconnect_btn = server.gui.add_button("Disconnect",
                                               icon=viser.Icon.PLUG_CONNECTED_X)
        mirror_cb = server.gui.add_checkbox(
            "Mirror motion to arm", initial_value=False,
            hint="Send the simulated joint angles to the ESP 50 times a second")
        waypoint_cb = server.gui.add_checkbox(
            "Waypoint mode for sequences", initial_value=True,
            hint="In a sequence, send each move once (W) and let the ESP plan the smooth "
                 "motion; the 3D arm then follows the encoders. Manual control still streams J.")
        traffic_md = server.gui.add_markdown("")
    with server.gui.add_folder("Message format (for the ESP firmware)",
                               expand_by_default=False):
        server.gui.add_markdown(
            "One text line per message, ending in a newline. Joint order: "
            + ", ".join(JOINT_NAMES).lower() + ". Angles in degrees; "
            "0 = arm stretched out horizontally, positive = up "
            "(shoulder, elbow, wrist).\n\n"
            "**PC to ESP**\n"
            "- `J,base,shoulder,elbow,wrist` joint angles (50 Hz while mirroring)\n"
            "- `G,1` / `G,0` close / open gripper\n"
            "- `W,base,shoulder,elbow,wrist,speed,id` waypoint move planned on the ESP "
            "(sequences in waypoint mode)\n"
            "- `H` halt a waypoint move and hold\n"
            "- `S` stop (software): hold still and ignore `J` until `R`\n"
            "- `R` reset after an emergency stop, software stop or fault\n\n"
            "**ESP to PC**\n"
            "- `A,base,shoulder,elbow,wrist` measured angles from the AS5600s "
            "(50 Hz)\n"
            "- `I,gripper` measured gripper servo current in amps (INA219)\n"
            "- `E,1` / `E,0` physical E-stop pressed / released\n"
            "- `K,1` grasp confirmed, `K,0` grasp failed (gripper re-opened), "
            "`K,2` grip lost\n"
            "- `F,joint,why` fault (ENC encoder, TRACK tracking): arm holds "
            "until `R`\n"
            "- any other line is shown here as the last message\n\n"
            "If no `J` arrives for 200 ms the ESP holds the last target.\n\n"
            "Example: `J,0.0,60.0,-70.0,-80.0`"
        )

# ---------------- Setup tab ----------------
with tabs.add_tab("Setup", icon=viser.Icon.SETTINGS):
    with server.gui.add_folder("Display"):
        dark_cb = server.gui.add_checkbox("Dark mode", initial_value=False)
        show_trail = server.gui.add_checkbox("Show tip trail", initial_value=True)
        show_presets = server.gui.add_dropdown(
            "Preset markers", ["Used in sequence", "All", "None"],
            initial_value="Used in sequence",
            hint="Which saved positions to show as blue markers")
        show_labels = server.gui.add_checkbox("Show step numbers", initial_value=True)
        show_grid = server.gui.add_checkbox("Show ground", initial_value=True)
        show_frames = server.gui.add_checkbox(
            "Show joint frames", initial_value=False,
            hint="Coordinate frame at each joint and at the gripping point "
                 "(red = x, green = y, blue = z). The pitch joints rotate about x.")
        cam_btn = server.gui.add_button("Reset camera", icon=viser.Icon.CAMERA)
    with server.gui.add_folder("Floor check"):
        server.gui.add_markdown(FLOOR_MSG)
    with server.gui.add_folder("Workspace map"):
        ws_height = server.gui.add_number("Scan height Z (mm)",
                                          initial_value=CUBE_SIZE_MM / 2, step=5, min=0)
        ws_scan = server.gui.add_button("Scan workspace", color="green",
                                        icon=viser.Icon.RADAR)
        ws_clear = server.gui.add_button("Clear map", icon=viser.Icon.TRASH)
        ws_progress = server.gui.add_progress_bar(0.0, color=BRAND)
        ws_md = server.gui.add_markdown(
            "🟩 gripper can point straight down  \n"
            "🟦 gripper tilted toward level (Auto mode)  \n"
            "🟨 reachable, but not at the chosen angle(s)  \n"
            "(no colour) out of reach")


# =============================================================================
# HELPERS
# =============================================================================
def set_fields(p):
    """Show a position (given in metres) in the Control tab fields, in mm."""
    x_in.value, y_in.value, z_in.value = (round(float(p[0]) * MM, 1),
                                          round(float(p[1]) * MM, 1),
                                          round(float(p[2]) * MM, 1))


def set_target(p):
    """Move the target. Ignored while the emergency stop is active."""
    if state["estop"]:
        return
    state["stopping"] = False
    p = np.array(p, dtype=float)
    p[2] = max(p[2], TIP_MIN_Z)
    target.position = tuple(p)
    set_fields(p)


def begin_stop():
    """Brake smoothly (within the acceleration limit), then hold that pose
    until the next command."""
    if not state["estop"]:
        state["stopping"] = True


def set_mode(mode):
    state["mode"] = mode
    icons = {"Manual": "🕹️", "Sequence": "📋"}
    mode_md.content = f"**Mode:** {icons[mode]} {mode}"


set_mode("Manual")


def blocked():
    """True (with a message) if the E-stop is active."""
    if state["estop"]:
        notify("Emergency stop active", "Reset the emergency stop first",
               color="red")
        return True
    return False


@target.on_update
def _(_):
    if state["estop"]:
        target.position = tuple(tip_position())  # snap back: arm is frozen
        return
    state["stopping"] = False
    set_fields(target.position)
    if state["mode"] != "Manual":
        stop_everything("Stopped: target moved by hand")


def stop_everything(message="Stopped"):
    was = state["mode"]
    if state.get("follow"):
        # waypoint mode: halt the ESP's move and take over from where the arm is
        state["follow"] = False
        hw_send("H")
        state["prev_qd"] = np.zeros(model.nv)
        target.position = tuple(tip_position())
        set_fields(target.position)
    set_mode("Manual")
    if was == "Sequence":
        seq_md.content = render_sequence(message)
    if cycle["active"]:
        finish_cycle(completed=False, note=message)


def do_grab():
    dist = np.linalg.norm(tip_position() - np.array(cube.position))
    if dist < 0.04:
        state["holding"] = True
        cube_md.content = "✋ **Held by gripper**"
        cube_label.text = f"{CUBE_MASS * 1000:.0f} g (held)"
        hw_send("G,1")
        return True
    cube_md.content = f"❌ Too far to grab ({dist * MM:.0f} mm away)"
    return False


def do_release():
    state["holding"] = False
    p = np.array(cube.position)
    cube.position = (float(p[0]), float(p[1]), CUBE_SIZE / 2)
    cube_md.content = "On the ground"
    cube_label.text = f"{CUBE_MASS * 1000:.0f} g"
    hw_send("G,0")


def refresh_preset_dropdowns():
    step_preset.options = list(PRESETS.keys()) + [CUSTOM]


def group_by_position(items, tol_mm=5.0):
    """Group (label, position) pairs that sit at (almost) the same spot.
    Returns a list of (position, [labels])."""
    groups = []
    for label, p in items:
        p = np.asarray(p, dtype=float)
        for g in groups:
            if np.linalg.norm(g[0] - p) * MM < tol_mm:
                if label not in g[1]:
                    g[1].append(label)
                break
        else:
            groups.append((p, [label]))
    return groups


def draw_presets():
    """Blue markers for presets. Names above the marker; presets at the same
    spot share one label. Which ones are shown: see the Setup tab."""
    for ball, _label in preset_nodes.values():
        ball.remove()  # removing the ball also removes its label
    preset_nodes.clear()
    mode = show_presets.value
    if mode == "None":
        return
    used = [s["pos"] for s in seq["steps"] if s["kind"] == "move"]
    items = []
    for name, p in PRESETS.items():
        if mode == "Used in sequence" and not any(
                np.linalg.norm(np.asarray(p) - u) * MM < 5.0 for u in used):
            continue
        items.append((name, p))
    for i, (p, names) in enumerate(group_by_position(items)):
        ball = server.scene.add_icosphere(
            f"/presets/p{i}", radius=0.01, color=(70, 130, 230), opacity=0.6,
            position=tuple(p),
        )
        label = server.scene.add_label(
            f"/presets/p{i}/label", " / ".join(names), position=(0.0, 0.0, 0.02),
            anchor="bottom-center",
        )
        preset_nodes[i] = (ball, label)


draw_presets()


# =============================================================================
# EMERGENCY STOP
# =============================================================================
@estop_btn.on_click
def _(_):
    trigger_estop()


def handle_hw_event(line):
    """Grasp results (K,1 / K,0 / K,2) and faults (F,joint,why) from the ESP."""
    if line == "K,1":
        notify("Grasp confirmed", "Gripper current above threshold", color="green")
    elif line in ("K,0", "K,2"):
        state["holding"] = False
        p = np.array(cube.position)
        cube.position = (float(p[0]), float(p[1]), CUBE_SIZE / 2)
        cube_label.text = f"{CUBE_MASS * 1000:.0f} g"
        what = "Grasp failed" if line == "K,0" else "Grip lost"
        cube_md.content = f"❌ {what} (gripper current)"
        stop_everything(f"❌ {what}: no gripper current")
        notify(what, "Sequence stopped", color="red")
    elif line.startswith("F,"):
        trigger_estop(source="fault", detail=line[2:])


def trigger_estop(source="button", detail=""):
    if state["estop"]:
        return
    stop_everything("🛑 Emergency stop")
    state["estop"] = True
    state["stopping"] = False
    state["prev_qd"] = np.zeros(model.nv)
    target.position = tuple(tip_position())
    hw_send("S")
    estop_btn.visible = False
    where = ("the **physical E-stop** on the arm" if source == "hardware"
             else f"an **arm fault** ({detail})" if source == "fault"
             else "the button")
    estop_banner.content = ("## 🛑 EMERGENCY STOP ACTIVE\n"
                            f"Triggered by {where}. The arm is frozen and "
                            "commands are blocked.")
    estop_banner.visible = True
    estop_reset.visible = True
    notify("EMERGENCY STOP", "All motion stopped", color="red")


@estop_reset.on_click
def _(_):
    if hw["estop_pressed"]:
        notify("Physical E-stop still pressed",
               "Twist to release the E-stop on the arm first", color="red")
        return
    state["estop"] = False
    set_fields(target.position)
    hw_send("R")
    estop_banner.visible = False
    estop_reset.visible = False
    estop_btn.visible = True
    notify("Emergency stop reset", "Motion commands are allowed again",
           color="green")


# =============================================================================
# CONTROL TAB ACTIONS
# =============================================================================
@go_btn.on_click
def _(_):
    if blocked():
        return
    p = np.array([x_in.value, y_in.value, z_in.value]) / MM
    if not validate_point(p, show_in_sequence_tab=False):
        return
    stop_everything()
    set_target(p)


@preset_save.on_click
def _(_):
    p = custom_point()
    if not validate_point(p):
        return
    name = preset_name.value.strip() or f"Point {len(PRESETS)}"
    if name == CUSTOM:
        name = "Custom point"
    PRESETS[name] = tuple(float(v) for v in p)
    refresh_preset_dropdowns()
    draw_presets()
    step_preset.value = name
    notify("Preset saved", f"'{name}' added to the list", color="green")


@grab_btn.on_click
def _(_):
    if not blocked():
        do_grab()


@release_btn.on_click
def _(_):
    if not blocked():
        do_release()


@place_btn.on_click
def _(_):
    if not state["holding"]:
        cube.position = (cube_x.value / MM, cube_y.value / MM, CUBE_SIZE / 2)


# =============================================================================
# SEQUENCE TAB
# =============================================================================
def describe(step):
    if step["kind"] == "move":
        return f"Move to **{step['name']}**"
    if step["kind"] == "grab":
        return "Grab cube"
    if step["kind"] == "release":
        return "Release cube"
    return f"Wait {step['secs']:.0f} s"


def render_sequence(note=""):
    steps = seq["steps"]
    if not steps:
        return "*No steps yet*"
    running = state["mode"] == "Sequence"
    lines = []
    for i, s in enumerate(steps):
        if running and i == seq["index"]:
            mark = "▶️"
        elif running and i < seq["index"]:
            mark = "✅"
        else:
            mark = f"{i + 1}."
        lines.append(f"{mark} {describe(s)}  ")
    if note:
        lines.append(f"\n{note}")
    return "\n".join(lines)


def draw_sequence():
    """Numbered markers in 3D for every move step, joined by a path line."""
    for ball, _label in seq_nodes:
        ball.remove()  # removing the ball also removes its label
    seq_nodes.clear()
    points = []
    numbered = []
    for i, s in enumerate(seq["steps"]):
        if s["kind"] != "move":
            continue
        points.append(np.asarray(s["pos"], dtype=float))
        numbered.append((str(i + 1), s["pos"]))
    # one marker per spot; step numbers BELOW it (preset names are above)
    for k, (p, nums) in enumerate(group_by_position(numbered)):
        ball = server.scene.add_icosphere(
            f"/seq_points/s{k}", radius=0.012, color=(255, 80, 80),
            position=tuple(p),
        )
        label = server.scene.add_label(
            f"/seq_points/s{k}/label", ", ".join(nums), position=(0.0, 0.0, -0.02),
            anchor="top-center", visible=show_labels.value,
        )
        seq_nodes.append((ball, label))
    if len(points) < 2:
        server.scene.add_line_segments(
            "/seq_path", points=np.zeros((0, 2, 3)), colors=(255, 80, 80),
            visible=False,
        )
        return
    # draw each move as the arc the tip actually follows
    segs = []
    for a, b in zip(points[:-1], points[1:]):
        n = max(2, int(arc_length(a, b) / 0.01))
        pts = [arc_point(a, b, t) for t in np.linspace(0.0, 1.0, n + 1)]
        segs += [[pts[j], pts[j + 1]] for j in range(n)]
    segs = np.array(segs)
    server.scene.add_line_segments(
        "/seq_path", points=segs, colors=(255, 80, 80),
        thickness=2.0, thickness_units="screen",
    )


def refresh_sequence(note=""):
    seq_md.content = render_sequence(note)
    draw_sequence()
    draw_presets()


def custom_point():
    """The Custom X, Y, Z fields, converted to metres."""
    return np.array([cx_in.value, cy_in.value, cz_in.value], dtype=float) / MM


def fmt_mm(p):
    return f"({p[0] * MM:.0f}, {p[1] * MM:.0f}, {p[2] * MM:.0f}) mm"


def update_step_fields(_=None):
    is_move = step_type.value == "Move to position"
    custom = is_move and step_preset.value == CUSTOM
    step_preset.visible = is_move
    for h in (cx_in, cy_in, cz_in, copy_btn):
        h.visible = custom
    if not custom:
        check_md.visible = False


step_type.on_update(update_step_fields)
step_preset.on_update(update_step_fields)


@copy_btn.on_click
def _(_):
    p = np.array(target.position) * MM
    cx_in.value, cy_in.value, cz_in.value = (round(float(p[0]), 1),
                                             round(float(p[1]), 1),
                                             round(float(max(p[2], 0.0)), 1))


def validate_point(p, show_in_sequence_tab=True):
    """Check a point (metres) before using it. Shows an error if it is bad."""
    ok, message = check_point(p)
    if show_in_sequence_tab:
        check_md.content = message
        check_md.visible = bool(message)
    if not ok:
        notify("Position out of reach", f"{fmt_mm(p)}: {message[2:]}", color="red")
    elif message:
        notify("Position check", f"{fmt_mm(p)}: {message[2:]}", color="yellow")
    return ok


def make_step(kind_label, preset):
    if kind_label == "Move to position":
        if preset == CUSTOM:
            p = custom_point()
            return {"kind": "move", "name": fmt_mm(p), "pos": p}
        return {"kind": "move", "name": preset, "pos": np.array(PRESETS[preset])}
    if kind_label == "Grab cube":
        return {"kind": "grab"}
    if kind_label == "Release cube":
        return {"kind": "release"}
    return {"kind": "wait", "secs": 1.0}


def editing_allowed():
    if state["mode"] == "Sequence":
        notify("Sequence running", "Stop it before editing", color="yellow")
        return False
    return True


@step_add.on_click
def _(_):
    if not editing_allowed():
        return
    new_step = make_step(step_type.value, step_preset.value)
    if new_step["kind"] == "move" and not validate_point(new_step["pos"]):
        return
    seq["steps"].append(new_step)
    refresh_sequence()


@step_undo.on_click
def _(_):
    if seq["steps"] and editing_allowed():
        seq["steps"].pop()
        refresh_sequence()


@step_clear.on_click
def _(_):
    if editing_allowed():
        seq["steps"].clear()
        seq_progress.value = 0.0
        trial_md.content = ""
        refresh_sequence()


@step_demo.on_click
def _(_):
    if not editing_allowed():
        return
    plan = ["Home", "Above pick", "Pick", "GRAB", "Above pick",
            "Above place", "Place", "RELEASE", "Above place", "Home"]
    seq["steps"] = []
    for p in plan:
        if p == "GRAB":
            seq["steps"].append({"kind": "grab"})
        elif p == "RELEASE":
            seq["steps"].append({"kind": "release"})
        elif p in PRESETS:
            seq["steps"].append(make_step("Move to position", p))
    refresh_sequence()


def start_step(i):
    seq["index"] = i
    seq["started"] = False
    seq["arrived_at"] = None
    n_steps = len(seq["steps"])
    done_steps = (seq["trial"] - 1) * n_steps + i
    seq_progress.value = 100.0 * done_steps / (seq["trials_total"] * n_steps)
    seq_md.content = render_sequence()


def start_trial():
    k, n = seq["trial"], seq["trials_total"]
    if reset_cube_cb.value and not state["holding"]:
        cube.position = (cube_x.value / MM, cube_y.value / MM, CUBE_SIZE / 2)
    start_cycle(f"Trial {k}/{n}" if n > 1 else "Sequence")
    trial_md.content = f"**Trial {k} of {n}**" if n > 1 else ""
    start_step(0)


@seq_run.on_click
def _(_):
    if blocked():
        return
    if not seq["steps"]:
        notify("Nothing to run", "Add some steps first", color="yellow")
        return
    stop_everything()
    seq["trial"] = 1
    seq["trials_total"] = max(1, int(trials_in.value))
    seq["first_report"] = len(reports)
    set_mode("Sequence")
    state["follow"] = waypoint_ready()
    if state["follow"]:
        notify("Waypoint mode", f"The ESP plans each move: joints at up to "
               f"{max(5.0, np.degrees(speed_cap())):.0f}°/s (base max 60°/s). "
               "Change it with Speed (%) in the Control tab.", color="blue")
    if waypoint_cb.value and mirror_cb.value and hw["mode"] == "serial" and not state["follow"]:
        notify("Waypoint mode not used", "No encoder data from the ESP: streaming J instead",
               color="yellow")
    start_trial()


@seq_stop.on_click
def _(_):
    stop_everything("⏹️ Stopped by user")
    begin_stop()


def abort_sequence(reason):
    if seq["trials_total"] > 1:
        reason += f" (trial {seq['trial']} of {seq['trials_total']})"
    stop_everything(f"❌ {reason}")
    begin_stop()
    notify("Sequence stopped", reason, color="red")


def summarise_trials():
    """Mean and spread over the trials of the run that just finished."""
    global trials_summary
    runs = [r for r in reports[seq["first_report"]:] if r["result"] == "Completed"]
    if len(runs) < 2:
        trials_summary = ""
        return
    t = np.array([r["duration_s"] for r in runs])
    peak = np.array([r["peak_total_sim"] for r in runs])
    energy = np.array([r["energy_sim_J"] for r in runs])
    trials_summary = (
        f"**Last run: {len(runs)} trials**  \n"
        f"Cycle time: {t.mean():.1f} ± {t.std(ddof=1):.1f} s  \n"
        f"Peak total current: {peak.mean():.2f} ± {peak.std(ddof=1):.2f} A  \n"
        f"Energy: {energy.mean():.0f} ± {energy.std(ddof=1):.0f} J"
    )
    report_hist_md.content = render_report_history()


def waypoint_ready():
    """Waypoint mode can be used: switched on, mirroring to a real ESP that sends angles."""
    return (waypoint_cb.value and mirror_cb.value and hw["mode"] == "serial"
            and angles_are_live() and not state["estop"])


def send_waypoint(p):
    """Work out the joint angles for target p and send them once as W,...,id.
    The ESP plans the smooth motion and replies Y,id when it has arrived."""
    angle = np.radians(target_angle_deg(p))
    q_goal = None
    for seed in seeds_for(p):
        ok, q = reach_test(p, seed.copy(), True, iters=200, d=check_data, angle=angle)
        if ok:
            q_goal = q
            break
    if q_goal is None:
        return False
    deg = np.degrees(q_goal[:NUM_JOINTS])
    # the real tool sits on the base axis line (see BASE_AIM_AT_TIP): aim the base at p
    deg[0] = np.degrees(np.arctan2(-p[1], p[0])) if BASE_AIM_AT_TIP else deg[0]
    hw["wp_id"] += 1
    seq["wp_id"] = hw["wp_id"]
    v = max(5.0, np.degrees(speed_cap()))
    hw_send("W," + ",".join(f"{a:.2f}" for a in deg) + f",{v:.1f},{hw['wp_id']}")
    return True


def follow_hardware():
    """Waypoint mode: the 3D arm follows the measured joint angles. The model's tool
    point is 7 mm to the side of the base axis, so its base is turned back by the same
    small angle the J stream adds (BASE_AIM_AT_TIP). Returns the joint speeds."""
    q_old = configuration.q.copy()
    q = q_old.copy()
    meas = hw["angles"]
    q[:NUM_JOINTS] = meas
    if BASE_AIM_AT_TIP:
        pin.framesForwardKinematics(model, check_data, q)
        t = check_data.oMf[TCP_ID].translation
        if np.hypot(t[0], t[1]) > 0.05:
            q[0] = meas[0] - (np.arctan2(-t[1], t[0]) - q[0])
    q = np.clip(q, model.lowerPositionLimit, model.upperPositionLimit)
    configuration.update(q)
    qd = np.zeros(model.nv)
    qd[:NUM_JOINTS] = (q[:NUM_JOINTS] - q_old[:NUM_JOINTS]) / STEP_DT[0]
    return qd


def update_sequence(now):
    steps = seq["steps"]
    i = seq["index"]
    s = steps[i]

    if not seq["started"]:
        seq["started"] = True
        seq["t0"] = now
        if s["kind"] == "move":
            set_target(s["pos"])
            if state.get("follow") and not send_waypoint(s["pos"]):
                abort_sequence(f"Step {i + 1}: '{s['name']}' has no joint solution")
            return
        if s["kind"] == "grab" and not do_grab():
            abort_sequence(f"Step {i + 1}: cube is not at the gripper")
            return
        if s["kind"] == "release":
            do_release()

    done = False
    if s["kind"] == "move" and state.get("follow"):
        if hw["wp_done"] == seq.get("wp_id"):
            if seq["arrived_at"] is None:
                seq["arrived_at"] = now
            done = now - seq["arrived_at"] > 0.3
        elif now - seq["t0"] > 15.0:
            abort_sequence(f"Step {i + 1}: the ESP did not finish the move")
            return
    elif s["kind"] == "move":
        if state["reach"] == "ok":
            if seq["arrived_at"] is None:
                seq["arrived_at"] = now
            done = now - seq["arrived_at"] > 0.3
        elif state["reach"] == "out":
            abort_sequence(f"Step {i + 1}: '{s['name']}' is out of reach")
            return
        elif now - seq["t0"] > 15.0:
            abort_sequence(f"Step {i + 1}: timed out")
            return
    elif s["kind"] == "wait":
        done = now - seq["t0"] > s["secs"]
    else:
        done = now - seq["t0"] > 0.4   # short pause after grab/release

    if not done:
        return
    if i + 1 < len(steps):
        start_step(i + 1)
        return

    finish_cycle(completed=True)
    if seq["trial"] < seq["trials_total"]:
        seq["trial"] += 1
        start_trial()
        return

    n = seq["trials_total"]
    seq_progress.value = 100.0
    summarise_trials()
    stop_everything("🏁 **Sequence finished**" + (f" ({n} trials)" if n > 1 else ""))
    notify("Sequence finished",
           (f"{n} trials completed. " if n > 1 else f"{len(steps)} steps completed. ")
           + "See the Report tab.", color="green")


# ---------------- Saved sequences ----------------
def list_sequences():
    if not SEQ_DIR.exists():
        return ["(none)"]
    names = sorted(p.stem for p in SEQ_DIR.glob("*.json"))
    return names or ["(none)"]


def clean_name(text, default):
    return "".join(c for c in text.strip() if c.isalnum() or c in "-_ ") or default


@seq_save_btn.on_click
def _(_):
    if not seq["steps"]:
        notify("Nothing to save", "The sequence is empty", color="yellow")
        return
    name = clean_name(seq_name.value, "sequence")
    SEQ_DIR.mkdir(exist_ok=True)
    path = SEQ_DIR / f"{name}.json"
    content = {
        "version": 2,
        "units": "mm",
        "saved": datetime.now().isoformat(timespec="seconds"),
        "presets_mm": {k: [round(v * MM, 1) for v in p] for k, p in PRESETS.items()
                       if k != "Home"},
        "cube_mm": [round(float(cube_x.value), 1), round(float(cube_y.value), 1)],
        "gripper_mode": "auto" if grip_mode.value == AUTO else "fixed",
        "gripper_deg": float(grip_angle.value),
        "steps": [step_to_json(s) for s in seq["steps"]],
    }
    path.write_text(json.dumps(content, indent=2), encoding="utf-8")
    seq_saved_dd.options = list_sequences()
    seq_saved_dd.value = name
    notify("Sequence saved", f"'{name}' ({len(seq['steps'])} steps)", color="green")


@seq_refresh_btn.on_click
def _(_):
    seq_saved_dd.options = list_sequences()


@seq_load_btn.on_click
def _(_):
    if not editing_allowed():
        return
    name = seq_saved_dd.value
    path = SEQ_DIR / f"{name}.json"
    if name == "(none)" or not path.exists():
        notify("Nothing to load", "Save a sequence first", color="yellow")
        return
    try:
        content = json.loads(path.read_text(encoding="utf-8"))
        steps = [step_from_json(d) for d in content["steps"]]
        presets = {k: tuple(v / MM for v in p)
                   for k, p in content.get("presets_mm", {}).items()}
        cube_start = content.get("cube_mm")      # where the cube starts (x, y) in mm
        g_mode = content.get("gripper_mode")     # "auto" or "fixed" (optional)
        g_deg = content.get("gripper_deg")
    except Exception as exc:
        notify("Could not load sequence", f"{path.name}: {exc}", color="red")
        return
    if presets:
        PRESETS.update(presets)   # adds the file's presets to your current list
        refresh_preset_dropdowns()
        draw_presets()
    if g_mode == "auto":
        grip_mode.value = AUTO
    elif g_mode == "fixed":
        grip_mode.value = FIXED
        if g_deg is not None:
            grip_angle.value = float(g_deg)
    grip_angle.visible = grip_mode.value == FIXED
    state["angle_key"] = None
    if cube_start is not None:
        cube_x.value, cube_y.value = float(cube_start[0]), float(cube_start[1])
        if not state["holding"]:
            cube.position = (cube_x.value / MM, cube_y.value / MM, CUBE_SIZE / 2)
    seq["steps"] = steps
    seq_progress.value = 0.0
    trial_md.content = ""
    seq_name.value = name
    refresh_sequence()
    notify("Sequence loaded", f"'{name}' ({len(steps)} steps)", color="green")


@seq_delete_btn.on_click
def _(_):
    name = seq_saved_dd.value
    path = SEQ_DIR / f"{name}.json"
    if name == "(none)" or not path.exists():
        return
    path.unlink()
    seq_saved_dd.options = list_sequences()
    notify("Sequence deleted", f"'{name}' removed")


seq_saved_dd.options = list_sequences()


# =============================================================================
# HARDWARE CONNECTION
# =============================================================================
def list_ports():
    ports = []
    if HAVE_SERIAL:
        ports = [p.device for p in serial.tools.list_ports.comports()]
    return [DEMO_PORT] + ports


hw_lock = threading.Lock()      # the main loop and the J sender thread both write


def hw_send(line):
    """Send one line to the ESP (does nothing unless connected to a real port)."""
    ser = hw["ser"]
    if hw["mode"] != "serial" or ser is None:
        return
    try:
        with hw_lock:
            ser.write((line + "\n").encode("ascii"))
        hw["sent"] += 1
    except Exception as exc:  # cable pulled, port closed, ...
        hw_disconnect(error=f"Write failed: {exc}")


def hw_reader(ser):
    """Background thread: read lines from the ESP."""
    while hw["ser"] is ser:
        try:
            raw = ser.readline()
        except Exception as exc:
            if hw["ser"] is ser:
                hw_disconnect(error=f"Read failed: {exc}")
            return
        if not raw:
            continue
        line = raw.decode("ascii", errors="ignore").strip()
        hw["received"] += 1
        hw["last_rx"] = line[:80]
        if line.startswith("#"):
            hw["info"] = line[1:].strip()[:160]
            continue
        if line.startswith("E,"):
            pressed = line[2:].strip() == "1"
            if pressed and not hw["estop_pressed"]:
                hw["estop_event"] = True     # handled in the main loop
            hw["estop_pressed"] = pressed
            continue
        if line.startswith("Y,"):                 # waypoint move finished on the ESP
            try:
                hw["wp_done"] = int(line[2:].split(",")[0])
            except ValueError:
                pass
            continue
        if line.startswith(("K,", "F,")):        # grasp result / arm fault
            hw["events"].append(line)
            continue
        if line.startswith(("I,", "A,")):
            try:
                values = np.array([float(v) for v in line[2:].split(",")])
            except ValueError:
                continue
            if line[0] == "I":                   # gripper current (A)
                if len(values) < 1:
                    continue
                hw["measured"] = float(values[0])
                hw["measured_time"] = time.time()
            else:                                # joint angles (deg)
                if len(values) != NUM_JOINTS:
                    continue
                hw["angles"] = np.radians(values)
                hw["angles_time"] = time.time()


def hw_connect(port, baud):
    hw_disconnect()
    hw["sent"] = hw["received"] = 0
    hw["last_rx"] = hw["error"] = ""
    if port == DEMO_PORT:
        hw["mode"] = "demo"
        hw["port"] = port
        return True
    if not HAVE_SERIAL:
        hw["error"] = "pyserial is not installed (run: pip install pyserial)"
        return False
    try:
        ser = serial.serial_for_url(port, baudrate=baud, timeout=0.1)
    except Exception as exc:
        hw["error"] = f"Could not open {port}: {exc}"
        return False
    hw["ser"] = ser
    hw["mode"] = "serial"
    hw["port"] = port
    threading.Thread(target=hw_reader, args=(ser,), daemon=True).start()
    hw_send("V,1")          # make sure the ESP streams A and I lines (may be off after testing)
    if state["estop"]:
        hw_send("S")
    return True


def hw_disconnect(error=""):
    ser = hw["ser"]
    hw["ser"] = None
    hw["mode"] = "offline"
    hw["measured"] = None
    hw["angles"] = None
    hw["estop_pressed"] = False
    if error:
        hw["error"] = error
        notify("Hardware disconnected", error, color="red")
    if ser is not None:
        try:
            ser.close()
        except Exception:
            pass


@port_refresh.on_click
def _(_):
    port_dd.options = list_ports()


@connect_btn.on_click
def _(_):
    if hw_connect(port_dd.value, int(baud_dd.value)):
        notify("Connected", hw["port"], color="green")
    else:
        notify("Connection failed", hw["error"], color="red")


@disconnect_btn.on_click
def _(_):
    hw_disconnect()
    mirror_cb.value = False


def measured_is_live():
    return (hw["mode"] == "serial" and hw["measured"] is not None
            and time.time() - hw["measured_time"] < MEASURED_TIMEOUT)


def render_hw_status():
    if hw["mode"] == "serial":
        light = f"🟢 **Connected** to `{hw['port']}`"
    elif hw["mode"] == "demo":
        light = "🟡 **Demo connection** (no real ESP)"
    else:
        light = "⚪ **Not connected**"
    if hw["error"]:
        light += f"  \n❌ {hw['error']}"
    if not HAVE_SERIAL:
        light += "  \n*pyserial not installed: only the demo connection is available.*"
    return light


def render_traffic():
    mirror = "on" if (mirror_cb.value and hw["mode"] != "offline") else "off"
    last = hw["last_rx"] or "-"
    return (f"Mirroring: **{mirror}**  \n"
            f"Sent: {hw['sent']}  |  Received: {hw['received']}  \n"
            f"Last message: `{last}`  \n"
            f"ESP info: `{hw['info'] or '-'}`")


port_dd.options = list_ports()


# =============================================================================
# CURRENTS: MODEL (ALL SERVOS) AND MEASURED (GRIPPER ONLY)
# =============================================================================
def simulated_currents(q, qd):
    tau = pin.computeGeneralizedGravity(model, data_dyn, q).copy()
    if state["holding"]:
        pin.computeJointJacobians(model, data_dyn, q)
        pin.updateFramePlacements(model, data_dyn)
        J = pin.getFrameJacobian(model, data_dyn, TCP_ID,
                                 pin.ReferenceFrame.LOCAL_WORLD_ALIGNED)
        tau += J[:3].T @ np.array([0.0, 0.0, CUBE_MASS * 9.81])
    tau = tau[:NUM_JOINTS]
    joints = IDLE_CURRENT + np.abs(tau) / KT + 0.25 * np.abs(qd[:NUM_JOINTS])
    grip = GRIP_HOLD_CURRENT if state["holding"] else GRIP_IDLE_CURRENT
    return tau, np.append(joints, grip)     # 4 joints + gripper (A)


def fictitious_measured(sim_grip):
    """Pretend INA219 reading of the gripper current: gain error, offset,
    first-order lag and noise. Replaced by real I messages from the ESP."""
    global fake_state
    alpha = DT / (FAKE_LAG + DT)
    fake_state = fake_state + alpha * (sim_grip * FAKE_GAIN + FAKE_OFFSET - fake_state)
    return max(0.0, fake_state + 0.02 * np.random.randn())


def angles_are_live():
    return (hw["mode"] == "serial" and hw["angles"] is not None
            and time.time() - hw["angles_time"] < MEASURED_TIMEOUT)


def fictitious_angles(cmd, tau):
    """Pretend AS5600 readings: servo lag, sag under load, calibration offset,
    and noise. Replaced by real A messages once the ESP sends them."""
    global fake_angles
    alpha = DT / (FAKE_ANGLE_LAG + DT)
    fake_angles = fake_angles + alpha * (cmd - fake_angles)
    sag = -np.radians(FAKE_SAG_DEG_PER_NM) * tau      # droops the way gravity pulls
    noise = np.radians(FAKE_ANGLE_NOISE_DEG) * np.random.randn(NUM_JOINTS)
    return fake_angles + sag + FAKE_ANGLE_OFFSET + noise


def render_angle_source():
    if angles_are_live():
        return "**Measured angles:** 🟢 live from the AS5600s"
    return ("**Measured angles:** 🟡 fictitious "
            "(no `A` messages from an ESP yet)")


def render_angles(cmd, meas):
    window = np.array(angle_window) if angle_window else None
    rows = ["| Joint | Commanded | Measured | Error | RMS error (5 s) |",
            "|---|---|---|---|---|"]
    for j in range(NUM_JOINTS):
        err = np.degrees(meas[j] - cmd[j])
        rms = (np.degrees(np.sqrt(np.mean((window[:, 1, j] - window[:, 0, j]) ** 2)))
               if window is not None else 0.0)
        rows.append(f"| {JOINT_NAMES[j]} | {np.degrees(cmd[j]):+.1f}° | "
                    f"{np.degrees(meas[j]):+.1f}° | {err:+.1f}° | {rms:.1f}° |")
    return "\n".join(rows)


def make_angle_figure(dark):
    fig = go.Figure()
    error_mode = angle_mode.value.startswith("Error")
    if angle_history:
        ts = np.array([h[0] for h in angle_history])
        ts = ts - ts[-1]
        cmd = np.degrees(np.array([h[1] for h in angle_history]))
        meas = np.degrees(np.array([h[2] for h in angle_history]))
        for j in selected_joints(angle_joint.value):
            colour, name = JOINT_COLOURS[j], JOINT_NAMES[j]
            if error_mode:
                fig.add_trace(go.Scatter(x=ts, y=meas[:, j] - cmd[:, j], mode="lines",
                                         name=name, line=dict(color=colour, width=2)))
                continue
            fig.add_trace(go.Scatter(x=ts, y=cmd[:, j], mode="lines",
                                     name=f"{name} cmd", legendgroup=name,
                                     line=dict(color=colour, width=2)))
            fig.add_trace(go.Scatter(x=ts, y=meas[:, j], mode="lines",
                                     name=f"{name} meas", legendgroup=name,
                                     line=dict(color=colour, width=1.5, dash="dot")))
    fig.update_layout(
        template="plotly_dark" if dark else "plotly_white",
        paper_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=45, r=10, t=10, b=40),
        xaxis_title="Time (s)",
        yaxis_title="Error (°)" if error_mode else "Angle (°)",
        legend=dict(orientation="h", y=-0.25),
    )
    if error_mode:
        fig.update_yaxes(range=[-6, 6])
    return fig


def selected_joints(choice):
    """Joint indices for a dropdown value ("All" or a joint name)."""
    return list(range(NUM_JOINTS)) if choice == "All" else [JOINT_NAMES.index(choice)]


def selected_currents(choice):
    """Current-channel indices for a dropdown value ("All" or a servo name)."""
    return list(range(len(CUR_NAMES))) if choice == "All" else [CUR_NAMES.index(choice)]


def make_figure(dark):
    fig = go.Figure()
    show = plot_mode.value
    g = len(CUR_NAMES) - 1                     # index of the gripper channel
    if history:
        ts = np.array([h[0] for h in history])
        ts = ts - ts[-1]
        sim = np.array([h[1] for h in history])
        meas = np.array([h[2] for h in history])
        if show == SHOW_MODEL:
            for j in selected_currents(plot_joint.value):
                fig.add_trace(go.Scatter(x=ts, y=sim[:, j], mode="lines",
                                         name=f"{CUR_NAMES[j]} model",
                                         line=dict(color=JOINT_COLOURS[j], width=2)))
        elif show == SHOW_GRIP:
            fig.add_trace(go.Scatter(x=ts, y=sim[:, g], mode="lines",
                                     name="Gripper model",
                                     line=dict(color=JOINT_COLOURS[g], width=2)))
            fig.add_trace(go.Scatter(x=ts, y=meas, mode="lines",
                                     name="Gripper measured",
                                     line=dict(color=JOINT_COLOURS[g], width=1.5,
                                               dash="dot")))
        else:
            fig.add_trace(go.Scatter(x=ts, y=meas - sim[:, g], mode="lines",
                                     name="Gripper error",
                                     line=dict(color=JOINT_COLOURS[g], width=2)))
    if show == SHOW_GRIP_ERR:
        y_title, y_range = "Error (A)", [-1.0, 1.0]
    else:
        chans = selected_currents(plot_joint.value) if show == SHOW_MODEL else [g]
        y_title = "Current (A)"
        y_range = [0, float(max(CUR_LIMIT[j] for j in chans)) * 1.2]
        if len(chans) == 1:   # one servo: show its stall current
            fig.add_hline(y=float(CUR_LIMIT[chans[0]]), line_dash="dash",
                          line_color="red", annotation_text="stall",
                          annotation_position="top left")
    fig.update_layout(
        template="plotly_dark" if dark else "plotly_white",
        paper_bgcolor="rgba(0,0,0,0)",
        margin=dict(l=45, r=10, t=10, b=40),
        xaxis_title="Time (s)", yaxis_title=y_title, yaxis_range=y_range,
        legend=dict(orientation="h", y=-0.25),
    )
    return fig


def render_source():
    if measured_is_live():
        return "**Measured gripper current:** 🟢 live from the INA219 (ESP)"
    return ("**Measured gripper current:** 🟡 fictitious "
            "(no `I` messages from an ESP yet)")


def render_currents(sim, meas, tau):
    """sim: model currents (4 joints + gripper); meas: measured gripper current."""
    rows = ["| Servo | Torque | Model | Measured | |",
            "|---|---|---|---|---|"]
    for j in range(len(CUR_NAMES)):
        is_grip = j == len(CUR_NAMES) - 1
        peak = max(sim[j], meas) if is_grip else sim[j]
        flag = ("🔴" if peak > CUR_LIMIT[j]
                else "🟡" if peak > 0.8 * CUR_LIMIT[j] else "🟢")
        torque = "-" if is_grip else f"{tau[j]:+.2f} Nm"
        measured = f"{meas:.2f} A" if is_grip else "not measured"
        rows.append(f"| {CUR_NAMES[j]} | {torque} | {sim[j]:.2f} A | "
                    f"{measured} | {flag} |")
    rows.append("")
    rows.append(f"**Total servo current (model):** {sim.sum():.2f} A  ")
    if compare_window:
        w = np.array(compare_window)
        rms = float(np.sqrt(np.mean((w[:, 1] - w[:, 0]) ** 2)))
        pct = 100.0 * (w[:, 1].mean() - w[:, 0].mean()) / max(w[:, 0].mean(), 1e-6)
        rows.append(f"**Gripper, measured vs model (5 s):** {pct:+.1f} %, "
                    f"RMS error {rms:.3f} A")
    return "\n".join(rows)


CSV_HEADER = (["time_s", "mode", "estop", "holding_cube",
               "target_x_mm", "target_y_mm", "target_z_mm",
               "tip_x_mm", "tip_y_mm", "tip_z_mm"]
              + [f"{n.lower()}_cmd_deg" for n in JOINT_NAMES]
              + [f"{n.lower()}_meas_deg" for n in JOINT_NAMES]
              + [f"{n.lower()}_torque_Nm" for n in JOINT_NAMES]
              + [f"{n.lower()}_I_model_A" for n in CUR_NAMES]
              + ["gripper_I_meas_A", "I_total_model_A", "gripper_angle_deg",
                 "angles_source", "gripper_current_source"])


def send_to_browser(event, path):
    if event.client is not None:
        event.client.send_file_download(path.name, path.read_bytes())


@clear_log_btn.on_click
def _(_):
    global log_start
    log_rows.clear()
    history.clear()
    compare_window.clear()
    angle_history.clear()
    angle_window.clear()
    log_start = time.time()
    log_txt.value = "0"


# =============================================================================
# CYCLE REPORT
# =============================================================================
def start_cycle(kind):
    cycle.update(active=True, kind=kind, t0=time.time(), samples=[])


def finish_cycle(completed, note=""):
    if not cycle["active"]:
        return
    cycle["active"] = False
    samples = cycle["samples"]
    if len(samples) < 2:
        return
    t = np.array([s[0] for s in samples])
    sim = np.array([s[1] for s in samples])
    meas = np.array([s[2] for s in samples])
    qd = np.array([s[3] for s in samples])
    held = np.array([s[4] for s in samples])
    duration = t[-1] - t[0]
    dt = np.diff(t, prepend=t[0])
    total_sim = sim.sum(axis=1)
    g = len(CUR_NAMES) - 1
    hold = held > 0.5
    grip_hold_sim = float(sim[hold, g].mean()) if hold.any() else float("nan")
    grip_hold_meas = float(meas[hold].mean()) if hold.any() else float("nan")

    report = {
        "number": len(reports) + 1,
        "kind": cycle["kind"],
        "result": "Completed" if completed else "Stopped",
        "note": note,
        "time": datetime.now().strftime("%H:%M:%S"),
        "duration_s": duration,
        "held_s": float((held * dt).sum()),
        "peak_sim": sim.max(axis=0),
        "rms_sim": np.sqrt((sim ** 2).mean(axis=0)),
        "max_speed": np.abs(qd).max(axis=0),
        "peak_total_sim": float(total_sim.max()),
        "avg_total_sim": float(total_sim.mean()),
        "energy_sim_J": float(((sim @ CUR_VOLTAGE) * dt).sum()),
        "charge_mAh": float((total_sim * dt).sum() / 3.6),
        "grip_peak_meas": float(meas.max()),
        "grip_hold_sim": grip_hold_sim,
        "grip_hold_meas": grip_hold_meas,
        "grip_error_pct": float(100.0 * (meas.mean() - sim[:, g].mean())
                                / max(sim[:, g].mean(), 1e-6)),
        "measured_source": "ESP" if measured_is_live() else "fictitious",
    }
    reports.append(report)
    report_md.content = render_report(report)
    report_hist_md.content = render_report_history()


def render_report(r):
    busiest = int(np.argmax(r["rms_sim"][:NUM_JOINTS]))
    hold = ("not held" if np.isnan(r["grip_hold_meas"]) else
            f"{r['grip_hold_meas']:.2f} A measured, {r['grip_hold_sim']:.2f} A model")
    lines = [
        f"#### Cycle {r['number']}: {r['kind']} ({r['result'].lower()})",
        f"Finished at {r['time']}  ",
        f"**Cycle time:** {r['duration_s']:.1f} s  ",
        f"**Cube carried for:** {r['held_s']:.1f} s  ",
        f"**Peak total current (model):** {r['peak_total_sim']:.2f} A  ",
        f"**Average total current (model):** {r['avg_total_sim']:.2f} A  ",
        f"**Energy used (servo rails, model):** {r['energy_sim_J']:.0f} J  ",
        f"**Charge used (model):** {r['charge_mAh']:.1f} mAh  ",
        f"**Gripper current while holding:** {hold}  ",
        f"**Gripper peak measured:** {r['grip_peak_meas']:.2f} A; "
        f"measured vs model {r['grip_error_pct']:+.1f} % "
        f"(measured data: {r['measured_source']})  ",
        f"**Busiest joint:** {JOINT_NAMES[busiest]}",
        "",
        "| Servo | Peak (model) | RMS (model) | Max speed |",
        "|---|---|---|---|",
    ]
    for j in range(len(CUR_NAMES)):
        speed = (f"{r['max_speed'][j]:.2f} rad/s" if j < NUM_JOINTS else "-")
        lines.append(f"| {CUR_NAMES[j]} | {r['peak_sim'][j]:.2f} A | "
                     f"{r['rms_sim'][j]:.2f} A | {speed} |")
    if r["note"]:
        lines.append(f"\n{r['note']}")
    return "\n".join(lines)


def render_report_history():
    if not reports:
        return "*None yet*"
    lines = ["| # | Type | Result | Time | Peak total | Energy |",
             "|---|---|---|---|---|---|"]
    for r in reports:
        lines.append(f"| {r['number']} | {r['kind']} | {r['result']} | "
                     f"{r['duration_s']:.1f} s | {r['peak_total_sim']:.2f} A | "
                     f"{r['energy_sim_J']:.0f} J |")
    if trials_summary:
        lines.append("\n" + trials_summary)
    return "\n".join(lines)


REPORT_COLUMNS = (
    ["Cycle", "Type", "Result", "Finished", "Cycle time (s)", "Cube carried (s)",
     "Peak total model (A)", "Avg total model (A)", "Energy model (J)",
     "Charge model (mAh)", "Gripper peak meas (A)", "Gripper holding meas (A)",
     "Gripper holding model (A)", "Gripper meas vs model (%)", "Measured data"]
    + [f"{n} peak model (A)" for n in CUR_NAMES]
    + [f"{n} RMS model (A)" for n in CUR_NAMES]
    + [f"{n} max speed (rad/s)" for n in JOINT_NAMES]
)


def nan_round(v, n=3):
    """Round, or an empty cell when the value is missing (NaN)."""
    return "" if np.isnan(v) else round(float(v), n)


def report_row(r):
    return ([r["number"], r["kind"], r["result"], r["time"],
             round(r["duration_s"], 2), round(r["held_s"], 2),
             round(r["peak_total_sim"], 3), round(r["avg_total_sim"], 3),
             round(r["energy_sim_J"], 1), round(r["charge_mAh"], 2),
             round(r["grip_peak_meas"], 3), nan_round(r["grip_hold_meas"]),
             nan_round(r["grip_hold_sim"]), round(r["grip_error_pct"], 1),
             r["measured_source"]]
            + [round(float(v), 3) for v in r["peak_sim"]]
            + [round(float(v), 3) for v in r["rms_sim"]]
            + [round(float(v), 3) for v in r["max_speed"]])


def write_sheet(ws_, header, rows, widths=14):
    from openpyxl.styles import Font, PatternFill
    ws_.append(header)
    for cell in ws_[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="E67828")
    for row in rows:
        ws_.append(row)
    ws_.freeze_panes = "A2"
    for i in range(1, len(header) + 1):
        ws_.column_dimensions[ws_.cell(1, i).column_letter].width = widths


XLSX_LOG_MAX = 150_000      # data-log rows put in the Excel file (50 min at 50 Hz); all rows always go to the CSV
DOWNLOAD_MAX_MB = 40        # larger files stay in the logs folder only
export_busy = {"on": False}


@export_all_btn.on_click
def _(event):
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        notify("Excel export needs openpyxl",
               "In the Miniforge Prompt run: pip install openpyxl", color="red")
        return
    if not log_rows and not reports:
        notify("Nothing to export", "Run the arm or a sequence first", color="yellow")
        return
    if export_busy["on"]:
        notify("Export already running", "Wait for the 'Export saved' message", color="yellow")
        return
    if state["mode"] == "Sequence":
        notify("Sequence running", "Export when the sequence has finished", color="yellow")
        return
    # Writing Excel keeps Python busy for seconds. The J stream is paused meanwhile, so the
    # ESP simply holds the arm (watchdog) instead of receiving late, bunched targets.
    export_busy["on"] = True
    rows, reps = list(log_rows), list(reports)          # snapshot; logging carries on
    notify("Exporting...", f"{len(reps)} cycles, {len(rows):,} log rows. The arm holds "
           "still until it is done.", color="blue")
    print(f"[export] started: {len(reps)} cycles, {len(rows)} log rows", flush=True)
    threading.Thread(target=export_worker, args=(event, rows, reps), daemon=True).start()


def export_worker(event, rows, reps):
    import traceback
    try:
        paths = export_files(rows, reps)
        notify("Export saved", " | ".join(str(p) for p in paths), color="green")
        print("[export] saved:", *paths, sep="\n  ", flush=True)
        for p in paths:
            if p.stat().st_size < DOWNLOAD_MAX_MB * 1e6:
                send_to_browser(event, p)
    except Exception as exc:
        traceback.print_exc()
        notify("Export failed", f"{type(exc).__name__}: {exc} (details in the terminal)", color="red")
    finally:
        export_busy["on"] = False


def export_files(log_rows, reports):
    """Excel file (summary, cycle reports, data log up to XLSX_LOG_MAX rows) + full data log as CSV."""
    from openpyxl import Workbook
    stamp = f"{datetime.now():%Y%m%d_%H%M%S}"
    LOG_DIR.mkdir(exist_ok=True)
    paths = []
    if log_rows:
        csv_path = LOG_DIR / f"arm_export_{stamp}_datalog.csv"
        with open(csv_path, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(CSV_HEADER)
            w.writerows(log_rows)
        paths.append(csv_path)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    info = [
        ["Exported", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Units", "positions in mm, angles in degrees, currents in A"],
        ["Angle convention", "0 = arm stretched out horizontally, positive = up"],
        ["Cycles in report", len(reports)],
        ["Data log samples", len(log_rows)],
        ["Joints", ", ".join(JOINT_NAMES)],
        ["Servo rail voltage (V)", ", ".join(f"{v:g}" for v in CUR_VOLTAGE)
         + "  (base, shoulder, elbow, wrist, gripper)"],
        ["Stall torque (Nm)", ", ".join(f"{v:g}" for v in STALL_TORQUE)],
        ["Torque constant Kt (Nm/A)", ", ".join(f"{v:.2f}" for v in KT)],
        ["Stall current per servo (A)", ", ".join(f"{v:g}" for v in CUR_LIMIT)],
        ["Current measurement", "gripper only (INA219); joint currents are model values"],
        ["Gripper angle", grip_mode.value if grip_mode.value == AUTO
         else f"Fixed {grip_angle.value:.0f} deg"],
        ["Cube mass (kg)", CUBE_MASS],
        ["Measured gripper current", "ESP" if measured_is_live() else "fictitious"],
        ["Measured angles", "ESP (AS5600)" if angles_are_live() else "fictitious"],
        [],
        ["Sequence steps"],
    ]
    info += [[f"{i + 1}. {describe(s).replace('**', '')}"]
             for i, s in enumerate(seq["steps"])]
    if trials_summary:
        info += [[], ["Trials summary"]]
        info += [[line.replace("**", "").strip()] for line in trials_summary.split("\n")]
    write_sheet(summary, ["Item", "Value"], info, widths=40)

    write_sheet(wb.create_sheet("Cycle reports"), REPORT_COLUMNS,
                [report_row(r) for r in reports])
    if len(log_rows) <= XLSX_LOG_MAX:
        write_sheet(wb.create_sheet("Data log"), CSV_HEADER, log_rows)
    else:
        summary.append([])
        summary.append([f"Data log has {len(log_rows):,} rows: too many for this file, "
                        f"see {paths[0].name} (same folder)"])
    path = LOG_DIR / f"arm_export_{stamp}.xlsx"
    wb.save(path)
    paths.insert(0, path)
    return paths


@report_clear.on_click
def _(_):
    global trials_summary
    trials_summary = ""
    reports.clear()
    report_md.content = "*No cycles run yet*"
    report_hist_md.content = render_report_history()


# =============================================================================
# WORKSPACE MAP
# =============================================================================
ws = {"running": False}
ws_data = model.createData()


def reach_test(p, q, at_angle, iters=60, d=None, angle=None):
    """Damped least-squares IK. Returns (reached, final q).
    at_angle: also require the gripper tilt to equal `angle` (rad)."""
    d = ws_data if d is None else d
    angle = np.radians(grip_angle.value) if angle is None else angle
    for _ in range(iters):
        pin.framesForwardKinematics(model, d, q)
        e_pos = p - d.oMf[TCP_ID].translation
        J = pin.computeFrameJacobian(model, d, q, TCP_ID,
                                     pin.ReferenceFrame.LOCAL_WORLD_ALIGNED)
        J = J[:3, :NUM_JOINTS]
        if at_angle:
            e_ang = angle - gripper_angle(q)
            if np.linalg.norm(e_pos) < 0.003 and abs(e_ang) < 0.03:
                return above_floor(d), q
            err = np.append(e_pos, e_ang)
            J = np.vstack([J, [0.0, 1.0, 1.0, 1.0]])   # tilt = shoulder+elbow+wrist
        else:
            if np.linalg.norm(e_pos) < 0.003:
                return above_floor(d), q
            err = e_pos
        step = J.T @ np.linalg.solve(J @ J.T + 1e-3 * np.eye(len(err)), err)
        dq = np.zeros(model.nv)
        dq[:NUM_JOINTS] = np.clip(step, -0.5, 0.5)
        q = np.clip(pin.integrate(model, q, dq),
                    model.lowerPositionLimit, model.upperPositionLimit)
    return False, q


check_data = model.createData()


def seeds_for(p):
    """Starting guesses for the IK: current pose, home pose, and home pose turned
    toward the point (the base limits make a turned seed useful)."""
    turned = q0.copy()
    turned[0] = np.clip(-np.arctan2(p[1], p[0]), LOWER[0], UPPER[0])
    return [configuration.q.copy(), q0.copy(), turned]


auto_cache = {}


def auto_angle(p):
    """Steepest gripper angle (deg) from AUTO_ANGLES_DEG at which the tip can
    reach p, or None if none of them works. Results are cached (2 mm grid)."""
    key = tuple(np.round(np.asarray(p) * 500).astype(int))
    if key in auto_cache:
        return auto_cache[key]
    d = model.createData()
    seeds = seeds_for(p)
    result = None
    for a in AUTO_ANGLES_DEG:
        if any(reach_test(p, s.copy(), True, iters=80, d=d, angle=np.radians(a))[0]
               for s in seeds):
            result = a
            break
    if len(auto_cache) > 5000:
        auto_cache.clear()
    auto_cache[key] = result
    return result


def target_angle_deg(p):
    """Gripper angle (deg) to use for target p, in the current mode."""
    if grip_mode.value == FIXED:
        return float(grip_angle.value)
    a = auto_angle(p)
    if a is None:   # not reachable at any listed angle: level reaches furthest
        return 0.0 if np.hypot(p[0], p[1]) > 0.36 else -90.0
    return float(a)


def loop_target_angle(tp):
    """target_angle_deg for the control loop. In Auto mode it is worked out at most
    every 5 loops (0.1 s) while the target is moving, so dragging stays smooth."""
    if grip_mode.value == FIXED:
        return float(grip_angle.value)
    key = tuple(np.round(tp * 500).astype(int))
    if key != state.get("angle_key") and (state.get("angle_key") is None
                                          or loop % 5 == 0):
        state["angle_key"] = key
        state["angle_deg"] = target_angle_deg(tp)
    return state["angle_deg"]


@grip_mode.on_update
def _(_):
    grip_angle.visible = grip_mode.value == FIXED
    state["angle_key"] = None     # work the angle out again in the new mode


def check_point(p):
    """Returns (ok, message). ok is False if the arm cannot reach p at all."""
    if p[2] < TIP_MIN_Z - 1e-9:
        return False, (f"❌ Too low: the tip must stay at least {TIP_MIN_Z_MM:.0f} mm "
                       "above the floor.")
    if beyond_reach(p, 0.01):
        return False, "❌ Outside the arm's working area (too far away)."
    seeds = seeds_for(p)
    if grip_mode.value == AUTO:
        a = auto_angle(p)
        if a == AUTO_ANGLES_DEG[0]:
            return True, ""
        if a is not None:
            return True, (f"ℹ️ Beyond the pointing-down range: the gripper "
                          f"will tilt to {a}° here.")
        fail_msg = "⚠️ Reachable, but only with the gripper tilted upward."
    else:
        for seed in seeds:
            if reach_test(p, seed, True, iters=100, d=check_data)[0]:
                return True, ""
        fail_msg = (f"⚠️ Reachable, but not with the gripper at "
                    f"{grip_angle.value:.0f}°.")
    for seed in seeds:
        if reach_test(p, seed, False, iters=100, d=check_data)[0]:
            return True, fail_msg
    return False, "❌ Outside the arm's working area."


SCAN_TILTS_DEG = [-75, -60, -45, -30, -15, 0]   # tried after -90 in Auto mode


def scan_workspace(height):
    global workspace_cloud
    auto = grip_mode.value == AUTO
    first = np.radians(-90.0 if auto else grip_angle.value)
    xs = np.arange(-WORKSPACE_RADIUS, WORKSPACE_RADIUS + 1e-9, WORKSPACE_STEP)
    points, colours = [], []
    n_down = n_tilt = n_pos = n_total = 0
    for row, x in enumerate(xs):
        seed = q0.copy()
        for y in xs:
            r = np.hypot(x, y)
            if r > WORKSPACE_RADIUS or r < 0.04:
                continue
            n_total += 1
            p = np.array([x, y, height])
            fresh = q0.copy()
            fresh[0] = np.clip(-np.arctan2(y, x), LOWER[0], UPPER[0])

            def try_angle(a):
                for s0 in (seed, fresh):
                    ok, q = reach_test(p, s0.copy(), True, angle=a)
                    if ok:
                        return q
                return None

            q = try_angle(first)
            if q is not None:
                n_down += 1
                seed = q
                points.append(p)
                colours.append((60, 190, 90))
                continue
            if auto:
                for a in SCAN_TILTS_DEG:
                    q = try_angle(np.radians(a))
                    if q is not None:
                        break
                if q is not None:
                    n_tilt += 1
                    seed = q
                    points.append(p)
                    colours.append((70, 130, 230))
                    continue
            ok, q = reach_test(p, seed, False)
            if ok:
                n_pos += 1
                seed = q
                points.append(p)
                colours.append((240, 190, 40))
        ws_progress.value = 100.0 * (row + 1) / len(xs)

    cell = WORKSPACE_STEP ** 2
    if points:
        workspace_cloud = server.scene.add_point_cloud(
            "/workspace_map", points=np.array(points, dtype=np.float32),
            colors=np.array(colours, dtype=np.uint8),
            point_size=WORKSPACE_STEP * 0.9, point_shape="square",
            point_shading="flat",
        )
    first_label = "pointing down" if auto else f"gripper at {grip_angle.value:.0f}°"
    lines = [f"Map at Z = {height * MM:.0f} mm  ",
             f"🟩 {first_label}: **{n_down * cell:.3f} m²**  "]
    if auto:
        lines.append(f"🟦 gripper tilted toward level: **{n_tilt * cell:.3f} m²**  ")
    lines += [f"🟨 position only: **{n_pos * cell:.3f} m²**  ",
              f"Unreachable: {(n_total - n_down - n_tilt - n_pos) * cell:.3f} m² "
              "of the scanned area  ",
              "*Self-collisions are not checked.*"]
    ws_md.content = "\n".join(lines)
    ws["running"] = False
    notify("Workspace scan finished", f"Height {height * MM:.0f} mm", color="green")


@ws_scan.on_click
def _(_):
    if ws["running"]:
        return
    ws["running"] = True
    ws_progress.value = 0.0
    ws_md.content = "Scanning..."
    threading.Thread(target=scan_workspace, args=(ws_height.value / MM,),
                     daemon=True).start()


@ws_clear.on_click
def _(_):
    global workspace_cloud
    if workspace_cloud is not None:
        workspace_cloud.remove()
        workspace_cloud = None
    ws_progress.value = 0.0


# =============================================================================
# SEQUENCE FILES (positions stored in mm)
# =============================================================================
def step_to_json(s):
    out = {"kind": s["kind"]}
    if s["kind"] == "move":
        out.update(name=s["name"], pos_mm=[round(float(v) * MM, 1) for v in s["pos"]])
    if s["kind"] == "wait":
        out["secs"] = s["secs"]
    return out


def step_from_json(d):
    if d["kind"] == "move":
        if "pos_mm" in d:
            pos = np.array(d["pos_mm"], dtype=float) / MM
        else:
            pos = np.array(d["pos"], dtype=float)    # older files, in metres
        return {"kind": "move", "name": d["name"], "pos": pos}
    if d["kind"] == "wait":
        return {"kind": "wait", "secs": float(d.get("secs", 1.0))}
    return {"kind": d["kind"]}


# =============================================================================
# SETUP TAB: DISPLAY
# =============================================================================
@dark_cb.on_update
def _(_):
    server.gui.configure_theme(
        dark_mode=dark_cb.value, brand_color=BRAND,
        control_width="large", show_share_button=False,
    )
    apply_scene_theme(dark_cb.value)


@show_presets.on_update
def _(_):
    draw_presets()


@show_labels.on_update
def _(_):
    for _ball, label in seq_nodes:
        label.visible = show_labels.value


@show_frames.on_update
def _(_):
    for _link, node in frame_nodes:
        node.visible = show_frames.value
    if show_frames.value:
        update_frames()


@show_grid.on_update
def _(_):
    ground.visible = show_grid.value


@show_trail.on_update
def _(_):
    if trail is not None:
        trail.visible = show_trail.value


@cam_btn.on_click
def _(event):
    if event.client is not None:
        reset_camera(event.client)


# =============================================================================
# MAIN LOOP
# =============================================================================
def bar(fraction, width=10):
    n = int(round(np.clip(fraction, 0.0, 1.0) * width))
    return "▰" * n + "▱" * (width - n)


def update_reachability(tip, tp, qd_actual):
    if np.linalg.norm(tp - state["last_target"]) > 1e-6:
        state["last_target"] = tp.copy()
        state["stall_time"] = 0.0
        state["warned_target"] = None

    err_mm = np.linalg.norm(tip - tp) * 1000
    too_far = beyond_reach(tp)
    stopped = np.max(np.abs(qd_actual[:NUM_JOINTS])) < 0.02
    state["stall_time"] = state["stall_time"] + DT if stopped else 0.0

    if err_mm < IN_POSITION_MM:
        state["reach"] = "ok"
    elif too_far or state["stall_time"] > 0.6:
        state["reach"] = "out"
    else:
        state["reach"] = "moving"

    if state["estop"]:
        colour, text = (230, 60, 60), "🛑 **Emergency stop active**"
    elif state["reach"] == "ok":
        colour, text = (60, 200, 90), f"🟢 **In position** ({err_mm:.0f} mm off)"
    elif state["reach"] == "moving":
        colour, text = (240, 190, 40), f"🟡 **Moving:** {err_mm:.0f} mm to go"
    else:
        why = "beyond the arm's reach" if too_far else "the arm cannot get closer"
        colour = (230, 60, 60)
        text = f"🔴 **Unreachable**: {why}. Closest: {err_mm:.0f} mm away."
    kind_changed = state.get("status_kind") != (state["estop"], state["reach"])
    if kind_changed or loop % 5 == 0:
        state["status_kind"] = (state["estop"], state["reach"])
        target_ball.color = colour
        reach_md.content = text
    if state["reach"] == "out" and not state["estop"]:
        if state["warned_target"] is None and state["mode"] == "Manual":
            state["warned_target"] = tp.copy()
            notify("Target unreachable", f"The closest the tip can get is "
                   f"{err_mm:.0f} mm away", color="red")
    return err_mm


def speed_cap():
    """Joint speed allowed right now: the Speed % of the max joint speed."""
    return max_speed.value * speed.value / 100.0


def limit_motion(qd_request):
    """Apply the speed and acceleration limits. Returns the velocity used."""
    qd = qd_request.copy()
    prev = state["prev_qd"]
    vmax, amax = speed_cap(), max_accel.value
    wanted_v = np.max(np.abs(qd[:NUM_JOINTS]))
    wanted_a = np.max(np.abs(qd[:NUM_JOINTS] - prev[:NUM_JOINTS])) / DT
    # each joint: the Speed setting, but never above its servo's rated speed
    allowed = np.minimum(vmax, VEL_LIMIT)
    ratio = np.max(np.abs(qd[:NUM_JOINTS]) / allowed)
    if ratio > 1.0:
        qd *= 1.0 / ratio   # slow all joints together: the tip keeps its path
    dq = qd - prev
    step = np.max(np.abs(dq[:NUM_JOINTS]))
    if step > amax * DT:
        qd = prev + dq * (amax * DT / step)
    return qd, wanted_v, wanted_a


def render_joints(q, qd):
    rows = ["| Joint | Angle | Speed | Range used |", "|---|---|---|---|"]
    near = []
    for j in range(NUM_JOINTS):
        frac = (q[j] - LOWER[j]) / (UPPER[j] - LOWER[j])
        margin = np.degrees(min(q[j] - LOWER[j], UPPER[j] - q[j]))
        flag = " ⚠️" if margin < 5 else ""
        if margin < 5:
            near.append(JOINT_NAMES[j])
        rows.append(f"| {JOINT_NAMES[j]} | {np.degrees(q[j]):+6.1f}° | "
                    f"{qd[j]:+.2f} rad/s | {bar(frac)}{flag} |")
    rows.append(f"\n**Gripper angle:** {np.degrees(gripper_angle(q)):+.0f}° "
                "(shoulder + elbow + wrist; -90° = pointing down)")
    if near:
        rows.append(f"\n⚠️ **Near joint limit:** {', '.join(near)}")
    rows.append("\n*0° = arm stretched out horizontally; positive shoulder, "
                "elbow and wrist angles lift the arm.*")
    return "\n".join(rows)


def render_limits(wanted_v, wanted_a):
    vcap, amax = speed_cap(), max_accel.value
    lines = [f"Speed setting: **{speed.value:.0f} %** of {max_speed.value:.2f} rad/s "
             f"= **{vcap:.2f} rad/s**  ",
             f"Requested speed: {wanted_v:.2f} rad/s "
             f"{'(slowed)' if wanted_v > vcap else ''}  ",
             f"Requested accel: {wanted_a:.1f} / {amax:.1f} rad/s² "
             f"{'(slowed)' if wanted_a > amax else ''}  ",
             "Servo rated speeds: " + ", ".join(
                 f"{n} {v:.1f}" for n, v in zip(JOINT_NAMES, VEL_LIMIT)) + " rad/s  ",
             "\n*The arm never moves faster than these limits.*"]
    return "\n".join(lines)


loop = 0


def step():
    """One control loop: IK, limits, hardware, currents, readouts, logging."""
    global loop, trail
    now = time.time()
    STEP_DT[0] = DT if loop == 0 else float(np.clip(now - step.last, DT, 0.1))
    step.last = now

    # --- Keep the ESP streaming A and I lines (it starts with them off, e.g. after a reset) ---
    if (hw["mode"] == "serial" and time.time() - hw["angles_time"] > 1.0
            and now - hw.get("v_sent", 0.0) > 1.0):
        hw_send("V,1")
        hw["v_sent"] = now

    # --- Physical E-stop on the arm (reported by the ESP) ---
    if hw["estop_event"]:
        hw["estop_event"] = False
        trigger_estop(source="hardware")
    while hw["events"]:
        handle_hw_event(hw["events"].pop(0))

    if not state["estop"] and state["mode"] == "Sequence":
        update_sequence(now)

    # --- Inverse kinematics toward the target (frozen during an E-stop) ---
    tp = np.array(target.position)
    tp[2] = max(tp[2], TIP_MIN_Z)
    if state.get("follow") and not state["estop"] and not angles_are_live():
        stop_everything("❌ Waypoint mode: no encoder data from the ESP")
        notify("Sequence stopped", "No encoder data from the ESP", color="red")
    if state["estop"]:
        qd = np.zeros(model.nv)
        wanted_v = wanted_a = 0.0
    elif state.get("follow"):
        qd = follow_hardware()
        wanted_v = wanted_a = 0.0
    elif state["stopping"]:
        prev = state["prev_qd"]
        fastest = np.max(np.abs(prev))
        brake = max_accel.value * DT
        # slow all joints together, so the tip keeps its direction while braking
        qd = prev * max(0.0, 1.0 - brake / fastest) if fastest > 0 else prev * 0.0
        wanted_v = wanted_a = 0.0
        advance(qd)
        target.position = tuple(tip_position())
        tp = np.array(target.position)
        if np.max(np.abs(prev)) > 0 and not np.any(qd):
            set_fields(tp)   # stopped: hold this pose until the next command
    else:
        aim = arc_carrot(tip_position(), tp)
        task.set_target(tilt_target(configuration.q, aim,
                                    np.radians(loop_target_angle(tp))))
        velocity = solve_ik(configuration, [task], DT, solver="quadprog",
                            damping=1e-3, barriers=floor_barriers,
                            safety_break=False)
        qd, wanted_v, wanted_a = limit_motion(velocity)
        advance(qd)
    state["prev_qd"] = qd
    q = configuration.q
    show_robot(q)
    tip = tip_position()

    err_mm = update_reachability(tip, tp, qd)
    target_ball.position = tuple(tp)

    if state["holding"]:
        cube.position = (float(tip[0]), float(tip[1]), float(tip[2]))

    # --- Hardware: send joint angles while mirroring ---
    if mirror_cb.value and hw["mode"] == "serial" and not state["estop"] and not state.get("follow"):
        hw["j_q"] = np.degrees(q[:NUM_JOINTS]).copy()    # sent by j_sender() at a steady 50 Hz
        hw["j_qd"] = np.degrees(qd[:NUM_JOINTS]).copy()
        r_xy = float(np.hypot(tip[0], tip[1]))
        if BASE_AIM_AT_TIP and r_xy > 0.05:
            # The URDF puts the tool point 7 mm to the side of the base axis; the real arm
            # has no such offset (8.6, 13 Oct: at base 0 the tool sat on the 0 deg line).
            # So point the real base straight at the model's tool point instead.
            hw["j_q"][0] = np.degrees(np.arctan2(-tip[1], tip[0]))
        hw["j_time"] = time.time()
    else:
        hw["j_q"] = None

    # --- Joint angles: commanded vs measured (AS5600) ---
    cmd = q[:NUM_JOINTS].copy()
    angles_live = angles_are_live()
    tau, sim = simulated_currents(q, qd)
    fake_q = fictitious_angles(cmd, tau)
    meas_q = hw["angles"].copy() if angles_live else fake_q
    angle_window.append((cmd, meas_q.copy()))
    if loop % 3 == 0:
        angle_history.append((now, cmd, meas_q.copy()))

    # --- Currents: model for every servo, measured for the gripper only ---
    fake = fictitious_measured(sim[-1])
    live = measured_is_live()
    meas = float(hw["measured"]) if live else fake
    compare_window.append((sim[-1], meas))
    if loop % 3 == 0:
        history.append((now, sim.copy(), meas))
    if cycle["active"]:
        cycle["samples"].append((now, sim.copy(), meas,
                                 qd[:NUM_JOINTS].copy(), float(state["holding"])))

    if loop % 5 == 0:
        trail_points.append(tip.copy())
    if loop % LOG_EVERY == 0:
        if recording.value:
            log_rows.append(
                [round(now - log_start, 3), state["mode"], int(state["estop"]),
                 int(state["holding"])]
                + [round(v * MM, 1) for v in tp] + [round(v * MM, 1) for v in tip]
                + [round(np.degrees(v), 2) for v in cmd]
                + [round(np.degrees(v), 2) for v in meas_q]
                + [round(v, 4) for v in tau]
                + [round(v, 4) for v in sim]
                + [round(meas, 4), round(float(sim.sum()), 4),
                   round(np.degrees(gripper_angle(q)), 1),
                   "ESP" if angles_live else "fictitious",
                   "ESP" if live else "fictitious"]
            )

    # --- Readouts (not every loop, to keep the browser smooth) ---
    if loop % 5 == 0:
        tip_txt.value = f"{tip[0] * MM:.0f}, {tip[1] * MM:.0f}, {tip[2] * MM:.0f}"
        grip_txt.value = (f"{np.degrees(gripper_angle(q)):.0f}°  (target "
                          f"{loop_target_angle(tp):.0f}°, "
                          f"{'auto' if grip_mode.value == AUTO else 'fixed'})")
        joints_md.content = render_joints(q, qd)
        limits_md.content = render_limits(wanted_v, wanted_a)
        currents_md.content = render_currents(sim, meas, tau)
        source_md.content = render_source()
        angle_source_md.content = render_angle_source()
        angles_md.content = render_angles(cmd, meas_q)
        log_txt.value = str(len(log_rows))
        hw_status_md.content = render_hw_status()
        traffic_md.content = render_traffic()
    if loop % 10 == 0 and len(trail_points) > 1:
        pts = np.array(trail_points)
        segs = np.stack([pts[:-1], pts[1:]], axis=1)
        trail = server.scene.add_line_segments(
            "/trail", points=segs, colors=(70, 130, 230),
            thickness=2.0, thickness_units="screen", visible=show_trail.value,
        )
    if show_frames.value and loop % 2 == 0:
        update_frames()
    # plots are slow to rebuild; while mirroring to the arm, redraw them once a second
    every = 50 if (mirror_cb.value and hw["mode"] == "serial") else 15
    if loop % every == 0:
        plot.figure = make_figure(dark_cb.value)
    if loop % every == every // 2:
        angle_plot.figure = make_angle_figure(dark_cb.value)

    loop += 1


def j_sender():
    """Send the latest joint targets at a steady 50 Hz, whatever the main loop is doing.
    Stops sending if the main loop has not produced a target for 1 s, so the ESP
    watchdog still catches a frozen GUI."""
    nxt = time.perf_counter()
    while True:
        nxt += DT
        q, qd, t0 = hw["j_q"], hw["j_qd"], hw["j_time"]
        age = time.time() - t0
        if q is not None and age < 1.0 and not export_busy["on"]:
            # carry the motion on through a short stall of the main loop (max 0.1 s)
            qs = q + qd * min(age, 0.1)
            hw_send("J," + ",".join(f"{v + 0.0:.2f}" for v in qs))
        time.sleep(max(0.0, nxt - time.perf_counter()))
        if time.perf_counter() - nxt > 0.1:      # fell far behind: restart the schedule
            nxt = time.perf_counter()


def main():
    threading.Thread(target=j_sender, daemon=True).start()
    show_robot(configuration.q)
    print("Open http://localhost:8080 in your browser")
    print("Press Ctrl+C here to stop")
    try:
        while True:
            start = time.time()
            step()
            time.sleep(max(0.0, DT - (time.time() - start)))
    except KeyboardInterrupt:
        pass
    finally:
        hw_send("S")          # leave a connected arm stopped
        hw_disconnect()


if __name__ == "__main__":
    main()
